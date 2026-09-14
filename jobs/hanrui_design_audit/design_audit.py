#!/usr/bin/env python
"""Hanrui Fang guide-library and pairing audit (evidence report only; SLURM only).

Produces the answer matrix (Q1-Q12), file inventory, workbook / scaffold / pairing /
FASTQ / guide-counting / sample-metadata audits, contradictions, questions for the
experimental team, Markdown + self-contained HTML reports and provenance.
No thresholds are changed and no perturbation analysis is run.
"""
from __future__ import annotations

import collections
import gzip
import hashlib
import html
import io
import json
import os
import re
import subprocess
import sys
import tarfile
from datetime import datetime
from itertools import islice
from pathlib import Path

import h5py
import numpy as np
import pandas as pd

DATA = Path("/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/data/260720_HANRUI_FANG_2_HUMAN_10X")
PARENT = DATA.parent
REPO = Path("/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline")
OUT = REPO / "results" / "Hanrui_Fang_library_design_audit"
WELLS = {"HF011A": "HF011AF", "HF011B": "HF011BF", "HF012A": "HF012AF", "HF012B": "HF012BF"}
N_SAMPLE_READS = 200_000        # per FASTQ file for structure statistics
N_CONSENSUS_READS = 2_000_000   # R1 reads for the scaffold consensus (one file per library)
ANCHORS = {"A": "GTTTAAGAGCTA", "C": "GTTTCAGAGCTA", "canonical_T": "GTTTTAGAGCTA"}
TSO = "TTTCTTATATGGG"
ILLUMINA_ADAPTER = "AGATCGGAAGAGC"  # TruSeq read-2 adapter start: marks the end of the sequenced insert
# Reference scaffold sequences supplied by the analyst from the literature for ORIENTATION ONLY (typed, not from a
# file in the dataset); identity to them is annotation, not experimental evidence.
KNOWN_SCAFFOLDS = {
    "SpCas9 canonical sgRNA scaffold (Jinek/Mali 2013)": "GTTTTAGAGCTAGAAATAGCAAGTTAAAATAAGGCTAGTCCGTTATCAACTTGAAAAAGTGGCACCGAGTCGGTGC",
    "SpCas9 flip+extension (F+E, Chen 2013) scaffold": "GTTTAAGAGCTATGCTGGAAACAGCATAGCAAGTTTAAATAAGGCTAGTCCGTTATCAACTTGAAAAAGTGGCACCGAGTCGGTGC",
    "cr1 / 'CR1' modified scaffold (Adamson 2016 dual-guide, variant)": "GTTTCAGAGCTATGCTGGAAACAGCATAGCAAGTTGAAATAAGGCTAGTCCGTTATCAACTTGAAAAAGTGGCACCGAGTCGGTGC",
}
KEYWORDS_PAIR = ["pair_id", "construct_id", "vector_id", "guide_A", "guide_C", "scaffold_A", "scaffold_C", "slot", "oligo_pair", "pair", "construct", "vector", "cassette", "position"]
KEYWORDS_MODALITY = ["CRISPRi", "CRISPRa", "dCas9", "KRAB", "VPR", "VP64", "knockout", "knock-out", "Cas9", "Cas12", "base edit", "deletion", "cutting", "nuclease", "interference", "activation"]
KEYWORDS_DESIGN = ["MOI", "multiplicity", "replicate", "biological", "technical", "condition", "treatment", "dual", "two guides", "paired", "cassette", "tRNA", "scaffold", "tracr", "hU6", "mU6", "H1", "7SK"]

inventory: list[dict] = []
evidence_extracts: list[dict] = []
job_id = os.environ.get("SLURM_JOB_ID", "not_in_slurm")


def log(*a):
    print(*a, flush=True)


def add_file(path: Path, category: str, inspected: str, note: str = ""):
    try:
        st = path.stat()
        inventory.append({"path": str(path), "category": category, "size_bytes": st.st_size, "mtime": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M"),
                          "inspected_how": inspected, "note": note})
    except FileNotFoundError:
        inventory.append({"path": str(path), "category": category, "size_bytes": -1, "mtime": "", "inspected_how": "MISSING", "note": note})


def extract(path, etype, text, note=""):
    evidence_extracts.append({"source": str(path), "evidence_type": etype, "extract": str(text)[:600], "note": note})


def sh(cmd, cwd=None):
    return subprocess.run(cmd, shell=True, capture_output=True, text=True, cwd=cwd).stdout


def grep_keywords(text: str, keywords, source, etype, max_hits=6):
    hits = []
    for kw in keywords:
        for m in re.finditer(re.escape(kw), text, flags=re.IGNORECASE):
            s = max(0, m.start() - 90); e = min(len(text), m.end() + 90)
            hits.append((kw, text[s:e].replace("\n", " ")))
            if len(hits) >= max_hits:
                break
        if len(hits) >= max_hits:
            break
    for kw, ctx in hits:
        extract(source, etype, f"[{kw}] ...{ctx}...")
    return hits


# =====================================================================================
# 1. inventory + USB search
# =====================================================================================
def do_inventory():
    log("== inventory")
    usb = sh(f"find {PARENT} {REPO} -maxdepth 8 -iname '*usb*' 2>/dev/null").split()
    usb_result = {"usb_paths_found": usb, "search_roots": [str(PARENT), str(REPO)],
                  "statement": ("No directory or file named 'USB'/'usb' (or containing 'usb') exists under the data parent directory or the pipeline repository."
                                if not usb else f"USB-named paths found: {usb}")}
    for p in sorted(PARENT.iterdir()):
        if p.is_file():
            add_file(p, "parent_data_dir_file", "listed only (unrelated dataset objects: other projects' h5ad)", "not part of the Hanrui Fang dataset")
    cats = []
    for p in sorted(DATA.rglob("*")):
        if not p.is_file():
            continue
        rel = str(p.relative_to(DATA))
        if rel.startswith("data_audit/"):
            cat = "previous_audit_report" if p.suffix in (".md", ".html", ".csv", ".json", ".tsv", ".txt", ".log") else "previous_audit_other"
            how = "read (text) / listed" if p.suffix in (".md", ".csv", ".json", ".tsv", ".txt", ".log") else "listed"
        elif p.name.endswith((".fastq.gz", ".fq.gz")):
            cat = "fastq_guide" if any(p.parts[i] in WELLS.values() for i in range(len(p.parts))) else "fastq_gex"
            how = f"first {N_SAMPLE_READS:,} reads parsed" + (f" (+{N_CONSENSUS_READS:,} R1 reads for scaffold consensus)" if "R1" in p.name and cat == "fastq_guide" and "L004" in p.name else "")
        elif p.suffix == ".bam" or p.suffix == ".bai":
            cat, how = "cellranger_bam", "listed only"
        elif p.name.endswith("_outs.tar"):
            cat, how = "cellranger_outs_tar", "member list read (tar -t)"
        elif p.name.endswith("_info.tar.gz"):
            cat, how = "cellranger_info_tar", "members read (_cmdline/_invocation/_versions/_log)"
        elif p.name == "web_summary.html" or p.name.endswith("_web_summary.html"):
            cat, how = "cellranger_web_summary", "embedded JSON parsed"
        elif p.suffix in (".xlsx", ".xls"):
            cat, how = "design_workbook", "all sheets parsed"
        elif p.name == "README.txt":
            cat, how = "cellranger_readme", "read"
        elif p.name == "checksum.md5":
            cat, how = "checksum", "read"
        else:
            cat, how = "other", "listed"
        add_file(p, cat, how)
        cats.append(cat)
    # feature reference / whitelist / sample sheet / FASTA searches
    searches = {
        "fasta": "-iname '*.fa' -o -iname '*.fasta' -o -iname '*.fa.gz' -o -iname '*.fasta.gz' -o -iname '*.fna'",
        "feature_reference": "-iname '*feature*ref*' -o -iname '*feature_reference*' -o -iname '*featureref*'",
        "whitelist": "-iname '*whitelist*' -o -iname '*737K*' -o -iname '*3M-*' -o -iname '*barcodes*.txt*'",
        "sample_sheet": "-iname '*sample*sheet*' -o -iname 'SampleSheet*' -o -iname '*metadata*' -o -iname '*manifest*'",
        "protocol_methods": "-iname '*protocol*' -o -iname '*method*' -o -iname '*.docx' -o -iname '*.pdf' -o -iname '*.pptx'",
    }
    found = {}
    for k, expr in searches.items():
        hits = sh(f"find {DATA} {PARENT} -maxdepth 8 -type f \\( {expr} \\) 2>/dev/null").split()
        hits = [h for h in hits if "/data_audit/" not in h]
        found[k] = hits
    # repository files that implement / configure counting and pairing
    for rel in ("src/perturbseq_pipeline/guide_counting.py", "src/perturbseq_pipeline/guide_design.py", "src/perturbseq_pipeline/guide_qc.py", "src/perturbseq_pipeline/dual_guides.py",
                "src/perturbseq_pipeline/guides.py", "src/perturbseq_pipeline/io.py", "config/hanrui_fang.yaml", "config/Hanrui_fang.yaml", "config/Hanrui_fang_dual_guide.yaml",
                "config/Hanrui_Fang_pair_guide_full.yaml", "config/Hanrui_fang_sample_manifest.csv", "jobs/run_hanrui_qc.sbatch", "jobs/hanrui_fang/03_count_guides.slurm",
                "jobs/hanrui_fang/count_guides_hanrui.py", "jobs/hanrui_pair_full/audit_and_reference.py"):
        add_file(REPO / rel, "pipeline_code_or_config", "read")
    for rel in ("results/outputs/hanrui_fang_qc/logs/run.log", "results/outputs/hanrui_fang_qc/logs/resolved_config.yaml", "results/outputs/hanrui_fang_qc/tables/guide_counting_per_file.tsv",
                "results/Hanrui_fang_dual_guide/logs/03_count_guides.log", "results/Hanrui_fang_dual_guide/guide_quantification_methods.md",
                "results/Hanrui_fang_dual_guide/tables/guide_read_retention_funnel.csv", "results/Hanrui_Fang_pair_guide_full/audit/pair_guide_reference.csv",
                "results/Hanrui_Fang_pair_guide_full/combined/logs/resolved_config.yaml", "results/Hanrui_Fang_pair_guide_full/combined/tables/pair_guide_qc_per_lane.csv",
                "results/Hanrui_Fang_pair_guide_full/combined/tables/pair_assignment_per_lane.csv", "results/Hanrui_Fang_pair_guide_full/combined/Hanrui_Fang_pair_guide_processed.h5ad"):
        add_file(REPO / rel, "previous_pipeline_output", "read (obs/tables) or listed")
    for w in WELLS:
        for rel in (f"results/Hanrui_fang_dual_guide/guide_counts/{w}/features.tsv.gz", f"results/Hanrui_fang_dual_guide/guide_counts/{w}/matrix.mtx.gz", f"results/Hanrui_fang_dual_guide/guide_counts/{w}/{w}_guide_counting_stats.json",
                    f"results/outputs/hanrui_fang_inputs/cellranger/{w}/filtered_feature_bc_matrix/barcodes.tsv.gz", f"results/outputs/hanrui_fang_inputs/cellranger/{w}/filtered_feature_bc_matrix/features.tsv.gz"):
            add_file(REPO / rel, "derived_count_matrix_or_staged_cellranger", "read")
    return usb_result, found, collections.Counter(cats)


# =====================================================================================
# 2. previous audit reports (text evidence, checked later against files)
# =====================================================================================
def do_previous_reports():
    log("== previous reports")
    texts = {}
    for p in sorted((DATA / "data_audit").rglob("*")):
        if p.is_file() and p.suffix in (".md", ".csv", ".json", ".txt", ".tsv") and "extracted" not in p.parts and "tar_listings" not in p.parts and p.stat().st_size < 3_000_000:
            texts[str(p)] = p.read_text(errors="ignore")
    for p in sorted((DATA / "data_audit").rglob("*.html")):
        texts[str(p)] = re.sub(r"<[^>]+>", " ", p.read_text(errors="ignore"))
    for rel in ("docs/hanrui_fang_core_final_report.md", "results/Hanrui_fang_dual_guide/Hanrui_fang_dual_guide_summary.md", "results/Hanrui_fang_dual_guide/guide_design_audit.md",
                "results/Hanrui_fang_dual_guide/pipeline_comparison/audit_report_summary.md", "results/Hanrui_Fang_pair_guide_full/audit/initial_audit.md"):
        p = REPO / rel
        if p.is_file():
            texts[str(p)] = p.read_text(errors="ignore")
            add_file(p, "previous_pipeline_report", "read (text)")
    all_text = "\n".join(texts.values())
    hits_mod = {kw: len(re.findall(re.escape(kw), all_text, flags=re.IGNORECASE)) for kw in KEYWORDS_MODALITY}
    hits_pair = {kw: len(re.findall(r"\b" + re.escape(kw) + r"\b", all_text, flags=re.IGNORECASE)) for kw in KEYWORDS_PAIR}
    for src, t in texts.items():
        grep_keywords(t, ["CRISPRi", "CRISPRa", "dCas9", "KRAB", "knockout", "modality", "biological replicate", "technical replicate", "MOI", "pair map", "pairing", "dual-guide", "dual guide"], src, "previous_report_text", max_hits=4)
    # statements previously made about the library (to be checked)
    claims = []
    for src, t in texts.items():
        for pat in (r"[^.\n]*(?:each|every) (?:vector|construct)[^.\n]*\.", r"[^.\n]*scaffold[- ]A[^.\n]*scaffold[- ]C[^.\n]*\.", r"[^.\n]*pair map[^.\n]*\.", r"[^.\n]*modality[^.\n]*\.", r"[^.\n]*replicate[^.\n]*unknown[^.\n]*\."):
            for m in re.finditer(pat, t, flags=re.IGNORECASE):
                claims.append({"source": src, "claim": m.group(0).strip()[:400]})
                if len(claims) > 400:
                    break
    return texts, hits_mod, hits_pair, claims


# =====================================================================================
# 3. web summaries + Cell Ranger info tars + outs tar listings + README
# =====================================================================================
def do_cellranger():
    log("== cellranger")
    rows = []
    for w in WELLS:
        ws = DATA / "web_summaries" / f"{w}_web_summary.html"
        rec = {"well": w, "web_summary": str(ws)}
        if ws.is_file():
            t = ws.read_text(errors="ignore")
            m = re.search(r"const data = (\{.*?\})\s*</script>", t, flags=re.S) or re.search(r"window\.__data__\s*=\s*(\{.*?\});", t, flags=re.S)
            j = None
            if m:
                try:
                    j = json.loads(m.group(1))
                except Exception:
                    j = None
            if j is None:
                # fall back: find a large JSON blob
                for mm in re.finditer(r"(\{\"summary\".*?\})\s*;?\s*</script>", t, flags=re.S):
                    try:
                        j = json.loads(mm.group(1)); break
                    except Exception:
                        continue
            rec["web_summary_json_parsed"] = j is not None
            if j is not None:
                flat = json.dumps(j)
                for key in ("chemistry", "pipeline_version", "reference_path", "transcriptome", "include_introns", "feature_reference", "feature_ref", "cellranger", "Chemistry", "Pipeline Version", "Feature Reference"):
                    for mm in re.finditer(r'"%s"\s*:\s*("[^"]*"|null|true|false|[\d.]+)' % re.escape(key), flat):
                        rec.setdefault(f"json_{key}", mm.group(1)[:120]); break
                # table rows of the summary (label/value pairs)
                for mm in re.finditer(r'\["([^"]{3,60})",\s*"([^"]{1,60})"\]', flat):
                    k2, v2 = mm.group(1), mm.group(2)
                    if any(x in k2 for x in ("Chemistry", "Sample ID", "Transcriptome", "Pipeline Version", "Include introns", "Estimated Number of Cells", "Number of Reads", "Feature Reference", "Sequencing Saturation", "Mean Reads per Cell")):
                        rec.setdefault(f"tbl_{k2}", v2)
                crispr = bool(re.search(r"CRISPR Guide Capture|crispr|Antibody Capture|Feature Barcode", flat, flags=re.IGNORECASE))
                rec["mentions_feature_barcode_or_crispr_library"] = crispr
                extract(ws, "cellranger_web_summary_json", {k: v for k, v in rec.items() if k.startswith(("tbl_", "json_"))})
        info = sorted((DATA / w).rglob("*_cellranger_count_info.tar.gz"))
        if info:
            try:
                with tarfile.open(info[0], "r:gz") as tf:
                    names = tf.getnames()
                    rec["info_tar_members"] = len(names)
                    for want in ("_cmdline", "_invocation", "_versions", "_sitecheck", "_tags", "_perf"):
                        for n in names:
                            if n.endswith(want):
                                f = tf.extractfile(n)
                                if f:
                                    txt = f.read(20000).decode(errors="ignore")
                                    rec[f"info{want}"] = txt.strip()[:1500]
                                    extract(info[0], "cellranger_info_tar_member", f"{n}: {txt.strip()[:500]}")
                                break
            except Exception as e:  # pragma: no cover
                rec["info_tar_error"] = repr(e)
        outs = sorted((DATA / w).rglob("*_cellranger_count_outs.tar"))
        if outs:
            lst = sh(f"tar -tf {outs[0]}").split()
            rec["outs_tar_members"] = len(lst)
            rec["outs_tar_member_names"] = ";".join(sorted(set(Path(x).name for x in lst)))[:800]
            rec["outs_has_crispr_analysis"] = any("crispr" in x.lower() or "feature_reference" in x.lower() or "protospacer" in x.lower() for x in lst)
        rd = sorted((DATA / w).rglob("README.txt"))
        if rd:
            rec["readme_text"] = rd[0].read_text(errors="ignore")[:1500]
            extract(rd[0], "cellranger_readme", rec["readme_text"][:500])
        rows.append(rec)
    return pd.DataFrame(rows)


# =====================================================================================
# 4. workbook
# =====================================================================================
def do_workbook():
    log("== workbook")
    xl = sorted((DATA / "web_summaries").glob("*.xls*")) + sorted(p for p in DATA.rglob("*.xls*") if "data_audit" not in p.parts and "web_summaries" not in p.parts)
    xl = xl[0]
    book = pd.read_excel(xl, sheet_name=None, header=None, dtype=object)
    rows = []
    guides = None
    for name, raw in book.items():
        raw_full = raw.copy()
        raw = raw.dropna(how="all").dropna(axis=1, how="all")
        header = [str(x) for x in raw.iloc[0].tolist()]
        body = raw.iloc[1:].reset_index(drop=True); body.columns = header
        rows.append({"sheet": name, "item": "__sheet__", "value": f"raw dims {raw_full.shape[0]}x{raw_full.shape[1]}; non-empty {raw.shape[0]}x{raw.shape[1]}; header row = {header}", "detail": ""})
        for c in header:
            col = body[c]
            rows.append({"sheet": name, "item": f"column:{c}", "value": f"dtype={pd.api.types.infer_dtype(col, skipna=True)}; non_null={int(col.notna().sum())}; unique={int(col.astype(str).nunique())}",
                         "detail": "examples: " + "|".join(col.dropna().astype(str).head(4).tolist()) + (" | KEYWORD MATCH: " + ",".join(k for k in KEYWORDS_PAIR + ["scaffold", "sample", "replicate", "well", "condition"] if k.lower() in c.lower()) if any(k.lower() in c.lower() for k in KEYWORDS_PAIR + ["scaffold", "sample", "replicate", "well", "condition"]) else "")})
        if guides is None:
            guides = body
    # sequence analysis on the guide sheet
    seq_col = next((c for c in guides.columns if str(c).lower() in ("seq", "sequence", "protospacer", "spacer")), None)
    gene_col = next((c for c in guides.columns if str(c).lower() in ("gene", "target", "target_gene")), None)
    seq = guides[seq_col].astype(str).str.strip().str.upper()
    lens = seq.str.len().value_counts().to_dict()
    dup = seq.duplicated(keep=False)
    labels = guides[gene_col].astype(str).str.strip()
    ntc_mask = labels.str.upper().str.replace(r"[-_ ]", "", regex=True).isin(["NOTARGET", "NONTARGETING", "NTC"])
    rows.append({"sheet": "Sheet1", "item": "n_guide_records", "value": len(guides), "detail": ""})
    rows.append({"sheet": "Sheet1", "item": "n_unique_sequences", "value": int(seq.nunique()), "detail": f"duplicates: {int(dup.sum())}"})
    rows.append({"sheet": "Sheet1", "item": "sequence_lengths", "value": json.dumps(lens), "detail": "all nucleotide strings: " + str(bool(seq.str.fullmatch(r"[ACGT]+").all()))})
    rows.append({"sheet": "Sheet1", "item": "n_target_labels", "value": int(labels.nunique()), "detail": "labels: " + ", ".join(sorted(labels.unique())[:60])})
    rows.append({"sheet": "Sheet1", "item": "n_non_targeting_records", "value": int(ntc_mask.sum()), "detail": f"label(s): {sorted(labels[ntc_mask].unique())}"})
    rows.append({"sheet": "Sheet1", "item": "guides_per_target_label", "value": json.dumps(labels.value_counts().to_dict()), "detail": ""})
    first_g = seq.str[0] == "G"
    bwg_col = next((c for c in guides.columns if "beginswithg" in str(c).lower().replace("_", "")), None)
    if bwg_col:
        bwg = guides[bwg_col].astype(str).str.lower().isin(["true", "1", "yes"])
        equiv = bool((bwg == first_g).all())
        rows.append({"sheet": "Sheet1", "item": f"test:{bwg_col} == (first base of {seq_col} is G)", "value": equiv, "detail": f"{int(bwg.sum())} True vs {int(first_g.sum())} sequences starting with G; mismatches: {int((bwg != first_g).sum())}. => derived sequence property, NOT a scaffold label"})
    eff_col = next((c for c in guides.columns if "efficacy" in str(c).lower() or "score" in str(c).lower()), None)
    if eff_col:
        e = pd.to_numeric(guides[eff_col], errors="coerce")
        rows.append({"sheet": "Sheet1", "item": f"column:{eff_col}:summary", "value": f"min {e.min():.3f} median {e.median():.3f} max {e.max():.3f}; NA {int(e.isna().sum())}", "detail": "numeric on-target efficacy prediction (derived design metric; not a scaffold or pair field)"})
    for kind, kws in (("scaffold-related columns", ["scaffold", "tracr", "backbone"]), ("pair-related columns", ["pair", "partner", "dual"]), ("construct-related columns", ["construct", "vector", "plasmid", "oligo", "cassette", "slot", "position"]), ("sample/replicate columns", ["sample", "replicate", "well", "condition", "lane", "batch"])):
        hits = [c for c in guides.columns if any(k in str(c).lower() for k in kws)]
        rows.append({"sheet": "Sheet1", "item": kind, "value": hits if hits else "NONE", "detail": ""})
    # first/second-base composition (informational)
    rows.append({"sheet": "Sheet1", "item": "first_base_composition", "value": json.dumps(seq.str[0].value_counts().to_dict()), "detail": ""})
    # cell-level comments / hidden sheets / defined names via openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xl, read_only=False, data_only=False)
        rows.append({"sheet": "__workbook__", "item": "sheetnames(openpyxl)", "value": wb.sheetnames, "detail": f"defined names: {list(wb.defined_names.keys()) if hasattr(wb.defined_names, 'keys') else wb.defined_names}"})
        for wsn in wb.sheetnames:
            ws = wb[wsn]
            comments = [(c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment]
            rows.append({"sheet": wsn, "item": "sheet_state_and_comments", "value": f"state={ws.sheet_state}; max_row={ws.max_row}; max_col={ws.max_column}", "detail": f"cell comments: {comments[:5] if comments else 'none'}"})
        props = wb.properties
        rows.append({"sheet": "__workbook__", "item": "properties", "value": f"creator={props.creator}; created={props.created}; modified={props.modified}; lastModifiedBy={props.lastModifiedBy}; title={props.title}", "detail": ""})
    except Exception as e:  # pragma: no cover
        rows.append({"sheet": "__workbook__", "item": "openpyxl", "value": f"error {e!r}", "detail": ""})
    md5 = hashlib.md5(xl.read_bytes()).hexdigest()
    rows.append({"sheet": "__workbook__", "item": "file", "value": str(xl), "detail": f"md5={md5}; size={xl.stat().st_size}"})
    wa = pd.DataFrame(rows)
    return xl, wa, guides, seq, labels, ntc_mask


# =====================================================================================
# 5. FASTQ read structure + scaffold consensus
# =====================================================================================
def read_seqs(path: Path, n: int):
    with gzip.open(path, "rb") as fh:
        for i, line in enumerate(fh):
            if i % 4 == 1:
                yield line.rstrip(b"\n")
                if i // 4 + 1 >= n:
                    break


def read_headers(path: Path, n: int = 3):
    out = []
    with gzip.open(path, "rt") as fh:
        for i, line in enumerate(fh):
            if i % 4 == 0:
                out.append(line.strip())
                if len(out) >= n:
                    break
    return out


def revcomp(s: str) -> str:
    return s.translate(str.maketrans("ACGTN", "TGCAN"))[::-1]


def do_fastq(gex_barcodes: dict):
    log("== fastq")
    rows, scaf_rows = [], []
    consensus = {}
    for w, lib in WELLS.items():
        for kind, sample in (("guide", lib), ("gex", w)):
            fdir = DATA / sample / "fastq"
            files = sorted(fdir.rglob("*_L004_*_001.fastq.gz"))
            for f in files:
                read = re.search(r"_(R1|R2|I1|I2)_001", f.name).group(1)
                seqs = list(read_seqs(f, N_SAMPLE_READS))
                lens = collections.Counter(len(s) for s in seqs)
                hdrs = read_headers(f)
                row = {"library_type": kind, "sample": sample, "well": w, "file": str(f), "read": read, "n_reads_sampled": len(seqs), "read_length_mode": lens.most_common(1)[0][0], "read_lengths_top": json.dumps(dict(lens.most_common(3))),
                       "header_example": hdrs[0][:120] if hdrs else ""}
                if read in ("R1", "R2"):
                    S = [s.decode() for s in seqs]
                    bc_hits = sum(1 for s in S if s[:16] in gex_barcodes[w])
                    row["frac_first16_in_gex_filtered_barcodes"] = bc_hits / max(len(S), 1)
                    # UMI randomness (positions 17-28): mean per-position max base fraction
                    if all(len(s) >= 28 for s in S[:1000]):
                        block = np.array([list(s[16:28]) for s in S[:20000] if len(s) >= 28])
                        maxfrac = [max(collections.Counter(block[:, i]).values()) / block.shape[0] for i in range(12)]
                        row["umi_pos17_28_mean_max_base_fraction"] = float(np.mean(maxfrac))
                    for name, anchor in ANCHORS.items():
                        pos = [s.find(anchor) for s in S]
                        hits = [p for p in pos if p >= 0]
                        rc_hits = sum(1 for s in S if revcomp(anchor) in s)
                        row[f"anchor_{name}_forward_frac"] = len(hits) / max(len(S), 1)
                        row[f"anchor_{name}_revcomp_frac"] = rc_hits / max(len(S), 1)
                        if hits:
                            c = collections.Counter(hits)
                            row[f"anchor_{name}_position_top"] = json.dumps(dict(c.most_common(4)))
                    tso_pos = [s.find(TSO) for s in S]
                    th = [p for p in tso_pos if p >= 0]
                    row["tso_forward_frac"] = len(th) / max(len(S), 1)
                    row["tso_position_top"] = json.dumps(dict(collections.Counter(th).most_common(4))) if th else ""
                    if th:
                        # G-run between TSO end and anchor: distance distribution (anchor - (tso+13) - 20)
                        gaps = []
                        for s in S[:50000]:
                            t = s.find(TSO)
                            if t < 0:
                                continue
                            a = min([p for p in (s.find(ANCHORS["A"]), s.find(ANCHORS["C"])) if p >= 0], default=-1)
                            if a > 0:
                                gaps.append(a - (t + len(TSO)) - 20)
                        row["tso_to_protospacer_gap_top"] = json.dumps(dict(collections.Counter(gaps).most_common(6))) if gaps else ""
                        # composition of the gap bases
                        gapbases = collections.Counter()
                        for s in S[:20000]:
                            t = s.find(TSO)
                            a = min([p for p in (s.find(ANCHORS["A"]), s.find(ANCHORS["C"])) if p >= 0], default=-1)
                            if t >= 0 and a > t + len(TSO) + 20:
                                gapbases.update(s[t + len(TSO):a - 20])
                        row["gap_base_composition"] = json.dumps(dict(gapbases.most_common()))
                    # poly-T / polyA presence (R2 of 5' libraries)
                    row["frac_polyT10"] = sum(1 for s in S if "T" * 10 in s) / max(len(S), 1)
                    row["frac_polyA10"] = sum(1 for s in S if "A" * 10 in s) / max(len(S), 1)
                    if read == "R1" and kind == "guide":
                        row["mean_N_fraction"] = float(np.mean([s.count("N") / max(len(s), 1) for s in S]))
                else:
                    S = [s.decode() for s in seqs]
                    row["top_index_sequences"] = json.dumps(dict(collections.Counter(S).most_common(3)))
                rows.append(row)
                log(f"   {f.name}: len {row['read_length_mode']}")
        # scaffold consensus from guide R1
        f = sorted((DATA / lib / "fastq").rglob("*_L004_R1_001.fastq.gz"))[0]
        per_class = {k: np.zeros((110, 5), dtype=np.int64) for k in ("A", "C")}
        n_class = collections.Counter()
        full_hits = collections.Counter()
        idx = {b: i for i, b in enumerate("ACGTN")}
        for sb in read_seqs(f, N_CONSENSUS_READS):
            s = sb.decode()
            for k in ("A", "C"):
                p = s.find(ANCHORS[k], 40)
                if p >= 20:
                    n_class[k] += 1
                    seg = s[p - 20:p + 90]
                    M = per_class[k]
                    for i, b in enumerate(seg[:110]):
                        M[i, idx.get(b, 4)] += 1
                    for name, sc in KNOWN_SCAFFOLDS.items():
                        if sc[:60] in s:
                            full_hits[(k, name)] += 1
                    break
        for k in ("A", "C"):
            M = per_class[k]
            tot = M.sum(axis=1)
            cons = "".join("ACGTN"[int(np.argmax(M[i]))] if tot[i] else "-" for i in range(110))
            agree = np.where(tot > 0, M.max(axis=1) / np.maximum(tot, 1), np.nan)
            cov = tot / max(tot[20], 1)
            scaffold_obs = cons[20:]  # from anchor start
            end = next((i for i in range(len(scaffold_obs)) if cov[20 + i] < 0.5), len(scaffold_obs))
            scaffold_obs = scaffold_obs[:end]
            adapter_at = scaffold_obs.find(ILLUMINA_ADAPTER)
            trimmed = scaffold_obs[:adapter_at] if adapter_at > 0 else scaffold_obs
            best = None
            for name, sc in KNOWN_SCAFFOLDS.items():
                L = min(len(sc), len(trimmed))
                mm = sum(a != b for a, b in zip(sc[:L], trimmed[:L]))
                if best is None or mm < best[1]:
                    best = (name, mm, L)
            ident = 1 - best[1] / best[2] if best and best[2] else float("nan")
            consensus[(w, k)] = {"well": w, "guide_library": lib, "scaffold_class": k, "anchor": ANCHORS[k], "n_reads_sampled": N_CONSENSUS_READS, "n_reads_with_anchor": int(n_class[k]),
                                 "observed_consensus_from_anchor_raw": scaffold_obs, "illumina_adapter_found_at_nt": adapter_at,
                                 "observed_scaffold_consensus_from_anchor": trimmed, "observed_scaffold_length_nt": len(trimmed), "observed_length_with_>=50pct_coverage": end,
                                 "mean_per_position_agreement_first_60": float(np.nanmean(agree[20:80])), "protospacer_region_mean_agreement": float(np.nanmean(agree[:20])),
                                 "closest_reference_scaffold(annotation_only)": best[0] if best else "", "mismatches_to_closest_over_L": f"{best[1]} / {best[2]}" if best else "",
                                 "identity_to_closest": round(ident, 3), "match_call": ("close variant (>=95 % identity over the sequenced part)" if ident >= 0.95 else "no close match in the analyst reference set"),
                                 "insert_ends_in_adapter": adapter_at > 0, "reads_with_first60_of_known_scaffolds": json.dumps({n: int(v) for (kk, n), v in full_hits.items() if kk == k}), "source_file": str(f)}
            scaf_rows.append(consensus[(w, k)])
    return pd.DataFrame(rows), pd.DataFrame(scaf_rows)


# =====================================================================================
# 6. guide count matrices / multiplicity / counting method
# =====================================================================================
def do_counts_and_method(gex_barcodes):
    log("== counts and method")
    rows = []
    stats_all = {}
    for w in WELLS:
        d = REPO / "results" / "Hanrui_fang_dual_guide" / "guide_counts" / w
        st = json.load(open(d / f"{w}_guide_counting_stats.json"))["sample"]
        stats_all[w] = st
        with gzip.open(d / "features.tsv.gz", "rt") as fh:
            feats = [l.rstrip("\n").split("\t") for l in fh if l.strip()]
        with gzip.open(d / "barcodes.tsv.gz", "rt") as fh:
            bcs = [l.strip() for l in fh if l.strip()]
        with gzip.open(d / "matrix.mtx.gz", "rt") as fh:
            hdr = [l for l in fh if not l.startswith("%")][0].split()
        rows.append({"well": w, "matrix_dir": str(d), "n_features": len(feats), "feature_type": collections.Counter(f[2] for f in feats).most_common(1)[0][0], "n_barcodes": len(bcs),
                     "barcodes_match_gex_filtered": len(set(bcs) & {b.split('-')[0] for b in gex_barcodes[w]}) == len(bcs), "mtx_header": " ".join(hdr),
                     "source": "perturbseq_pipeline.guide_counting from guide FASTQ (SLURM 20045256)", "reads_total": st["reads_total"], "reads_with_scaffold_anchor": st["reads_with_scaffold_anchor"],
                     "reads_spacer_matched": st["reads_spacer_matched"], "reads_via_1mismatch": st.get("reads_spacer_matched_via_mismatch"), "reads_via_shift": st.get("reads_spacer_matched_via_shift"),
                     "reads_matched_barcode_in_gex": st["reads_matched_barcode_in_gex"], "reads_matched_barcode_not_in_gex": st["reads_matched_barcode_not_in_gex"], "unique_cell_guide_umis": st["unique_cell_guide_umis"],
                     "reads_scaffold_A": st["reads_scaffold_A"], "reads_scaffold_C": st["reads_scaffold_C"]})
    cm = pd.DataFrame(rows)
    # multiplicity from the pair-guide-full combined processed object (obs only)
    h5 = REPO / "results" / "Hanrui_Fang_pair_guide_full" / "combined" / "Hanrui_Fang_pair_guide_processed.h5ad"
    mult = {}
    if h5.is_file():
        import anndata as ad
        with h5py.File(h5, "r") as f:
            obs = ad.io.read_elem(f["obs"]) if hasattr(ad, "io") else ad.experimental.read_elem(f["obs"])
            uns_keys = list(f["uns"].keys())
            G = ad.io.read_elem(f["obsm/guide_counts"]) if hasattr(ad, "io") else ad.experimental.read_elem(f["obsm/guide_counts"])
            scaf = np.array([str(x) for x in (ad.io.read_elem(f["uns/guide_scaffolds"]) if hasattr(ad, "io") else ad.experimental.read_elem(f["uns/guide_scaffolds"]))]) if "guide_scaffolds" in uns_keys else None
        from scipy import sparse
        G = sparse.csr_matrix(G)
        if scaf is None:
            # the pair-guide-full object predates the uns['guide_scaffolds'] fix for scaffold_column: auto;
            # take the per-guide class from the pair reference used by that run (same guide order)
            with h5py.File(h5, "r") as f:
                names = [str(x) for x in (ad.io.read_elem(f["uns/guide_names"]) if hasattr(ad, "io") else ad.experimental.read_elem(f["uns/guide_names"]))]
            ref = pd.read_csv(REPO / "results/Hanrui_Fang_pair_guide_full/audit/pair_guide_reference.csv", dtype=str, keep_default_na=False).set_index("guide_id")
            scaf = ref.reindex(names)["scaffold"].fillna("unknown").to_numpy().astype(str)
            mult_note = "scaffold classes taken from audit/pair_guide_reference.csv (object lacks uns['guide_scaffolds']; io round-trip bug fixed in this commit)"
        else:
            mult_note = "scaffold classes from uns['guide_scaffolds']"
        n3 = np.asarray((G >= 3).sum(axis=1)).ravel()
        n1 = np.asarray((G >= 1).sum(axis=1)).ravel()
        tot = np.asarray(G.sum(axis=1)).ravel()
        nA = np.asarray((G[:, scaf == "A"] >= 3).sum(axis=1)).ravel() if scaf is not None else None
        nC = np.asarray((G[:, scaf == "C"] >= 3).sum(axis=1)).ravel() if scaf is not None else None
        top = np.asarray(G.max(axis=1).todense()).ravel()
        # depth-relative detection (>= 2 % of the cell's top guide and >= 3 UMIs)
        rel = np.asarray(((G >= 3).multiply(G >= (0.02 * top)[:, None])).sum(axis=1)).ravel()
        mult = {"object": str(h5), "scaffold_source_note": mult_note, "n_cells_qc_pass": int(G.shape[0]), "median_guide_umis_per_cell": float(np.median(tot)), "mean_guides_per_cell_ge3umi": float(n3.mean()), "median_guides_per_cell_ge3umi": float(np.median(n3)),
                "mean_guides_per_cell_ge1umi": float(n1.mean()), "frac_cells_gt1_guide_ge3umi": float((n3 > 1).mean()), "frac_cells_exactly_2_guides_ge3umi": float((n3 == 2).mean()),
                "frac_cells_ge1A_and_ge1C_ge3umi": float(((nA >= 1) & (nC >= 1)).mean()) if nA is not None else None, "frac_cells_exactly_1A_1C_ge3umi": float(((nA == 1) & (nC == 1)).mean()) if nA is not None else None,
                "median_guides_per_cell_depth_relative(>=3umi_and_>=2pct_of_top)": float(np.median(rel)), "frac_cells_exactly_2_guides_depth_relative": float((rel == 2).mean()),
                "frac_cells_both_slots_resolved(pair statuses)": float(obs["pair_assignment_status"].astype(str).isin(["pair_targeting", "pair_non_targeting", "pair_targeting_plus_ntc", "dual_target_ambiguous"]).mean()) if "pair_assignment_status" in obs else None,
                "pair_status_counts": obs["pair_assignment_status"].astype(str).value_counts().to_dict() if "pair_assignment_status" in obs else {}}
        add_file(h5, "previous_pipeline_output", "obs + obsm['guide_counts'] read for multiplicity statistics")
    # counting method evidence from code / config / logs
    gc_src = (REPO / "src/perturbseq_pipeline/guide_counting.py").read_text()
    doc = gc_src.split('"""')[1] if '"""' in gc_src else ""
    cfg_qc = (REPO / "config/hanrui_fang.yaml").read_text()
    fq_block = cfg_qc[cfg_qc.find("  fastq:"):cfg_qc.find("  multiplet:")]
    resolved = (REPO / "results/outputs/hanrui_fang_qc/logs/resolved_config.yaml").read_text() if (REPO / "results/outputs/hanrui_fang_qc/logs/resolved_config.yaml").is_file() else ""
    log03 = (REPO / "results/Hanrui_fang_dual_guide/logs/03_count_guides.log").read_text() if (REPO / "results/Hanrui_fang_dual_guide/logs/03_count_guides.log").is_file() else ""
    cmd03 = (REPO / "jobs/hanrui_fang/03_count_guides.slurm").read_text() if (REPO / "jobs/hanrui_fang/03_count_guides.slurm").is_file() else ""
    version = sh("git describe --always --dirty 2>/dev/null || git rev-parse --short HEAD", cwd=REPO).strip()
    pkg_version = sh("python -c 'import perturbseq_pipeline as p; print(getattr(p, \"__version__\", \"n/a\"))'", cwd=REPO).strip()
    method = [
        ("software", "perturbseq_pipeline.guide_counting (in-house streaming Python counter; not Cell Ranger, not kallisto/CRISPResso)", "src/perturbseq_pipeline/guide_counting.py", "code"),
        ("version", f"package {pkg_version}; git {version}; counting job commit 4799562 (jobs/hanrui_fang/03_count_guides.slurm, SLURM 20045256); earlier exact-only run SLURM 20045135 (config/hanrui_fang.yaml)", "git", "code/log"),
        ("exact_command", re.search(r"python jobs/hanrui_fang/count_guides_hanrui.py.*?(?=\n\S|\Z)", cmd03, flags=re.S).group(0).replace("\\\n", " ").strip() if cmd03 else "n/a", "jobs/hanrui_fang/03_count_guides.slurm", "job script"),
        ("guide_reference", "design_out_v2.xlsx column 'seq' (560 protospacers, 20 nt), guide ids synthesised '<gene>_<n>' (guide_design.load_guide_design); all 560 kept as matrix columns", "config/hanrui_fang.yaml guides.design", "config"),
        ("whitelist", "Cell Ranger FILTERED GEX barcodes of the paired well (restrict_to_gex_barcodes: true); no 10x whitelist file was used; reads whose 16-mer is not a called cell are counted but not stored", "config/hanrui_fang.yaml guides.fastq.restrict_to_gex_barcodes", "config/code"),
        ("barcode_correction", "NONE (exact 16-mer match to filtered barcodes; no 1-mismatch barcode rescue)", "guide_counting.py count_fastq_file: barcode_index.get(seq[:B])", "code"),
        ("umi_handling", "12-nt UMI (positions 17-28) packed base-4; exact UMI de-duplication per (cell, guide) via np.unique; no UMI error correction; reads with non-ACGT UMI dropped", "guide_counting.py (_UMI_TABLE, codes)", "code"),
        ("mismatch_tolerance", "protospacer: exact 20-mer 5' of the scaffold anchor; +/-1 nt positional shift; unambiguous 1-mismatch index (max_mismatches=1 in the 20045256 run; 0 in the 20045135 run). Scaffold anchor: exact 12-mer regex, no mismatches. Variants shared by two designed guides are dropped from the index.", "guide_counting.py build_protospacer_index/_resolve_guide; jobs/hanrui_fang/count_guides_hanrui.py --max-mismatches 1", "code/job"),
        ("error_correction", "none beyond the unambiguous 1-mismatch protospacer index", "guide_counting.py", "code"),
        ("multiple_guide_matches", "impossible by construction: designed protospacers are >= 3 mismatches apart (audit table), first exact match wins; ambiguous 1-mismatch variants removed", "results/Hanrui_fang_dual_guide/tables/guide_design_near_identical_protospacers.csv (empty)", "derived table"),
        ("scaffold_A_vs_C", "read-level: regex alternation (GTTTAAGAGCTA)|(GTTTCAGAGCTA) immediately 3' of the protospacer; per-guide class = majority anchor with purity >= 0.9 and >= 20 reads (guide_qc.infer_scaffold_classes); the class is EMPIRICAL, the workbook has no scaffold column", "guide_counting.py GuideReadSpec.anchor_regex; guide_qc.py", "code"),
        ("count_source", "guide FASTQ R1 (HF0xxxF libraries) -> UMI count matrices; Cell Ranger was run GEX-only (no CRISPR feature reference), so no Cell Ranger guide matrix exists", "web summaries (feature_reference null); outs tar listing", "cellranger output"),
        ("pair_assignment_evidence_basis", "the pair assignment operates on the derived UMI count matrix (pipeline output), which itself is computed directly from raw guide FASTQ reads by the counter above; no independent raw-read validation of pairs (e.g. read-level co-occurrence on one molecule) exists or is possible in single-read 5' data", "dual_guides.py; guide_counting.py", "code"),
    ]
    gm = pd.DataFrame(method, columns=["item", "value", "evidence_path", "evidence_type"])
    return cm, mult, gm, stats_all, doc, fq_block, log03


# =====================================================================================
# 7. sample metadata
# =====================================================================================
def do_samples(cr: pd.DataFrame, fq: pd.DataFrame):
    log("== samples")
    rows = []
    man = REPO / "config/Hanrui_fang_sample_manifest.csv"
    mdf = pd.read_csv(man) if man.is_file() else pd.DataFrame()
    for w, lib in WELLS.items():
        gex_files = sorted((DATA / w / "fastq").rglob("*.fastq.gz"))
        gfiles = sorted((DATA / lib / "fastq").rglob("*.fastq.gz"))
        lanes = sorted({re.search(r"_L(\d{3})_", f.name).group(1) for f in gex_files})
        sidx = sorted({re.search(r"_S(\d+)_", f.name).group(1) for f in gex_files})
        gsidx = sorted({re.search(r"_S(\d+)_", f.name).group(1) for f in gfiles})
        flow = sorted({f.parent.name for f in gex_files})
        idx_rows = fq[(fq.well == w) & (fq.read.isin(["I1", "I2"]))]
        crr = cr[cr.well == w].iloc[0].to_dict() if len(cr) else {}
        rows.append({"sample": w, "guide_library": lib, "gex_fastq_dir_type": "SC_5PGEX_10X", "guide_fastq_dir_type": "FB_OTHER_10X", "flowcell_run": ";".join(flow), "sequencing_lanes": ";".join(lanes),
                     "gex_sample_index_number": ";".join(sidx), "guide_sample_index_number": ";".join(gsidx),
                     "index_reads_top": "; ".join(f"{r.read}:{r.top_index_sequences[:80]}" for r in idx_rows.itertuples() if r.library_type == "gex"),
                     "cellranger_sample_id": crr.get("tbl_Sample ID", ""), "cellranger_chemistry": crr.get("tbl_Chemistry", crr.get("json_chemistry", "")),
                     "n_cells_cellranger": crr.get("tbl_Estimated Number of Cells", ""),
                     "manifest_condition": mdf.set_index("lane_id").loc[w, "condition"] if len(mdf) else "", "manifest_replicate": mdf.set_index("lane_id").loc[w, "replicate"] if len(mdf) else "",
                     "manifest_replicate_type": mdf.set_index("lane_id").loc[w, "replicate_type"] if len(mdf) else "", "manifest_condition_status": mdf.set_index("lane_id").loc[w, "condition_status"] if len(mdf) else "",
                     "manifest_origin": "config/Hanrui_fang_sample_manifest.csv written by the analyst (previous iteration) from the well ids; NOT supplied by the experimental team",
                     "classification": "separate 10x capture / GEM well (own Cell Ranger run, own barcode space, own guide-capture library) = ESTABLISHED; biological vs technical replicate and the meaning of HF011 vs HF012 = UNKNOWN (no experimental metadata file exists)",
                     "readme_text": crr.get("readme_text", "")[:300]})
    return pd.DataFrame(rows), mdf


# =====================================================================================
# 8. answers
# =====================================================================================
def build_answers(ctx):
    A = []
    def q(qid, question, status, answer, evidence_path, evidence_type, exact, confidence, uncertainty, needs_confirmation):
        A.append({"question_id": qid, "question": question, "status": status, "answer": answer, "evidence_path": evidence_path, "evidence_type": evidence_type,
                  "exact_relevant_item": exact, "confidence": confidence, "remaining_uncertainty": uncertainty, "experimental_confirmation_required": needs_confirmation})
    sc = ctx["scaffold_df"]; fq = ctx["fq"]; wa = ctx["wa"]; mult = ctx["mult"]; cr = ctx["cr"]; hits_mod = ctx["hits_mod"]
    gA = fq[(fq.library_type == "guide") & (fq.read == "R1")]
    a_frac = float(gA["anchor_A_forward_frac"].mean()); c_frac = float(gA["anchor_C_forward_frac"].mean()); t_frac = float(gA["anchor_canonical_T_forward_frac"].mean())
    both_class_cells = mult.get("frac_cells_ge1A_and_ge1C_ge3umi")
    one_one = mult.get("frac_cells_exactly_1A_1C_ge3umi")
    consA = sc[sc.scaffold_class == "A"].iloc[0]; consC = sc[sc.scaffold_class == "C"].iloc[0]
    ref = pd.read_csv(REPO / "results/Hanrui_Fang_pair_guide_full/audit/pair_guide_reference.csv", dtype=str, keep_default_na=False)
    ref_obs = ref[ref.scaffold.isin(["A", "C"])]
    purity = pd.to_numeric(ref_obs.scaffold_purity, errors="coerce")
    frac_pure = float((purity >= 0.99).mean()) if len(purity) else float("nan")
    ntc_ref = ref[ref.is_non_targeting.str.lower() == "true"]
    ntc_split = ntc_ref.scaffold.value_counts().to_dict()
    per_t = ref[ref.is_non_targeting.str.lower() != "true"].groupby("target_gene_name").agg(nA=("scaffold", lambda x: int((x == "A").sum())), nC=("scaffold", lambda x: int((x == "C").sum())))
    gap = json.loads(gA.iloc[0]["gap_base_composition"]) if gA.iloc[0].get("gap_base_composition") else {}
    gap_tot = sum(gap.values()) or 1
    gap_desc = ", ".join(f"{b} {100 * n / gap_tot:.0f} %" for b, n in sorted(gap.items(), key=lambda kv: -kv[1]))
    # Q1
    q("Q1", "Library architecture: are scaffold A and scaffold C physically intended to form one paired perturbation in the same cell?",
      "PARTIALLY_ANSWERED",
      f"Two distinct scaffold-associated sequence classes are directly observed in the guide reads (anchors GTTTAAGAGCTA in {100 * a_frac:.1f} % and GTTTCAGAGCTA in {100 * c_frac:.1f} % of sampled R1 reads; {100 * frac_pure:.1f} % of the observed designed guides are followed by one class in >= 99 % of their reads), and both classes are present in the same cell in "
      f"{100 * both_class_cells:.1f} % of QC-pass cells (>= 3 UMIs). This is consistent with, but does NOT prove, a dual-cassette construct: the same observation would arise from two independent single-guide libraries co-transduced at high MOI, or from one library with two scaffold variants. "
      "No file (workbook, README, Cell Ranger metadata, web summary, previous protocol) states the construct architecture. The 'A'/'C' labels themselves are a computational classification introduced by the pipeline.",
      f"{consA['source_file']}; {ctx['xl']}; {DATA}/web_summaries/*_web_summary.html",
      "FASTQ read consensus (direct) + workbook parse (direct) + absence of documentation",
      f"observed scaffold consensus A: {consA['observed_scaffold_consensus_from_anchor'][:50]}...; C: {consC['observed_scaffold_consensus_from_anchor'][:50]}...; workbook columns: gene, seq, On-Target Efficacy Score, BeginsWithG",
      "high for 'two classes exist'; low for 'physically paired constructs'",
      "whether A and C guides are on one vector (dual-cassette), on two vectors, or represent one library with mixed scaffolds; which promoter drives each cassette", True)
    # Q2
    q("Q2", "Guide multiplicity: was one A+C pair per cell expected, or was high MOI / multiple-guide delivery intentional?",
      "NOT_ANSWERED",
      f"No document states the intended MOI. Observed (combined object, 161,594 QC-pass cells, >= 3 UMIs): mean {mult.get('mean_guides_per_cell_ge3umi', float('nan')):.2f} / median {mult.get('median_guides_per_cell_ge3umi', float('nan')):.0f} detected guides per cell, "
      f"{100 * mult.get('frac_cells_gt1_guide_ge3umi', float('nan')):.1f} % of cells with > 1 guide, {100 * both_class_cells:.1f} % with both scaffold classes, only {100 * one_one:.1f} % with exactly one A and one C guide; median guide UMIs per cell {mult.get('median_guide_umis_per_cell', float('nan')):.0f}. "
      f"With a depth-relative detection rule (>= 3 UMIs and >= 2 % of the cell's top guide) the median is {mult.get('median_guides_per_cell_depth_relative(>=3umi_and_>=2pct_of_top)', float('nan')):.0f} guides and {100 * mult.get('frac_cells_exactly_2_guides_depth_relative', float('nan')):.1f} % of cells carry exactly two. "
      "Whether the excess reflects intentional high MOI, ambient/background guide RNA at this sequencing depth, or multiplets cannot be decided from the files.",
      mult.get("object", ""), "derived count matrix statistics (computational) + absence of protocol", f"pair_status_counts: {json.dumps(mult.get('pair_status_counts', {}))[:300]}",
      "high for the observed numbers; none for the expectation", "intended MOI / infection titre; whether one construct per cell was the design goal; expected background guide UMI level", True)
    # Q3
    q("Q3", "Exact guide pairing: does any file explicitly identify which A guide is paired with which C guide?",
      "NOT_ANSWERED",
      "No. The workbook has exactly four columns (gene, seq, On-Target Efficacy Score, BeginsWithG); none of pair_id / construct_id / vector_id / guide_A / guide_C / scaffold_A / scaffold_C / slot / oligo_pair occurs in any dataset file, Cell Ranger output, README or web summary. "
      "No file states that every same-target A+C combination is a valid construct. The pipeline's 'same-target pair' rule is therefore an inference. " f"Design geometry from reads: per target a median of {per_t.nA.median():.0f} scaffold-A and {per_t.nC.median():.0f} scaffold-C observed guides, i.e. a median of {int((per_t.nA * per_t.nC).median())} possible A x C combinations per target, of which the true constructs (if any) are unknown.",
      f"{ctx['xl']}; {REPO}/results/Hanrui_Fang_pair_guide_full/audit/pair_guide_reference.csv", "workbook parse + keyword search over all text files (direct)",
      f"workbook header: {wa[wa.item == '__sheet__'].value.iloc[0][:160]}; keyword hits in dataset text files: {json.dumps(ctx['hits_pair_files'])[:200]}",
      "high (absence established across all files)", "the A-C construct map (or a statement that combinations are random / all valid)", True)
    # Q4
    bw = wa[wa.item.str.startswith("test:")]
    q("Q4", "Excel design content: what does the workbook contain and are any columns merely derived sequence properties?",
      "ANSWERED",
      f"One worksheet 'Sheet1' (creator/lastModifiedBy and dates recorded in workbook_audit.csv), 560 records x 4 columns; 560 unique 20-nt sequences (0 duplicates), 46 target labels (45 genes incl. three '(rsID)' variants + 'NO-TARGET'), 127 NO-TARGET records; no scaffold, pair, construct, sample or replicate columns; "
      f"'BeginsWithG' is exactly the boolean 'first base of seq == G' ({bw.value.iloc[0] if len(bw) else 'n/a'}: {bw.detail.iloc[0][:120] if len(bw) else ''}) - a derived sequence property, not a scaffold label; 'On-Target Efficacy Score' is a numeric design prediction.",
      str(ctx["xl"]), "workbook parse (direct; openpyxl + pandas)", "columns: gene | seq | On-Target Efficacy Score | BeginsWithG", "high", "who produced the workbook and whether it is the final ordered library (creator/modified metadata recorded in workbook_audit.csv)", False)
    # Q5
    q("Q5", "Scaffold sequence evidence: which scaffold / constant motifs exist, from which source, full scaffold or anchor, documented or assumed?",
      "PARTIALLY_ANSWERED",
      f"Two scaffold-associated sequence classes are supported. Directly from reads (2 M R1 reads per library): class A consensus '{consA['observed_scaffold_consensus_from_anchor']}' ({consA['n_reads_with_anchor']:,} reads; closest known: {consA['closest_known_scaffold']}, {consA['mismatches_to_closest_over_L']} mismatches) and class C consensus '{consC['observed_scaffold_consensus_from_anchor']}' ({consC['n_reads_with_anchor']:,} reads; closest known: {consC['closest_known_scaffold']}, {consC['mismatches_to_closest_over_L']}). "
      f"Both consensuses run into the Illumina adapter after {consA['observed_scaffold_length_nt']} nt, i.e. the sequenced insert ends inside the scaffold (consistent with a library primer located in the scaffold; inference). Class A is a close variant of the F+E scaffold (identity {consA['identity_to_closest']}); class C has {consC['match_call']} (identity {consC['identity_to_closest']} to the nearest analyst-supplied reference) - the reference set is annotation only. "
      f"The 12-nt anchors used by the counter (config guides.fastq.scaffolds) are computational choices matching the first 12 nt of these consensuses; the canonical 'GTTTTAGAGCTA' anchor occurs in {100 * t_frac:.2f} % of reads. No FASTA, feature-reference or documentation file provides the scaffold sequences; no mismatch tolerance is applied to anchors. The observed scaffolds are read-derived, NOT experimentally documented.",
      f"{consA['source_file']}; {REPO}/config/hanrui_fang.yaml (guides.fastq.scaffolds)", "FASTQ consensus (direct) vs configuration (assumption)", f"A anchor {ANCHORS['A']} in {100 * a_frac:.1f} % of R1; C anchor {ANCHORS['C']} in {100 * c_frac:.1f} % of R1 (200k-read samples)",
      "high for the sequences; none for their experimental provenance (which plasmid / vendor scaffold)", "vendor / plasmid identity of the two scaffold variants; whether they encode different tracr structures on purpose (e.g. to avoid recombination in a dual cassette)", True)
    # Q6
    g1 = gA.iloc[0]; g2 = fq[(fq.library_type == "guide") & (fq.read == "R2")].iloc[0]; x1 = fq[(fq.library_type == "gex") & (fq.read == "R1")].iloc[0]
    q("Q6", "FASTQ read structure: where are cell barcode, UMI, guide, scaffold, sample index; orientation; read lengths?",
      "PARTIALLY_ANSWERED",
      f"Established from the reads themselves: guide-library R1 is {g1['read_length_mode']} nt; positions 1-16 match Cell Ranger filtered GEX barcodes in {100 * g1['frac_first16_in_gex_filtered_barcodes']:.1f} % of sampled reads (cell barcode); positions 17-28 are near-random (UMI; mean max-base fraction {g1.get('umi_pos17_28_mean_max_base_fraction', float('nan')):.2f}); "
      f"the 5' TSO '{TSO}' occurs at position {g1['tso_position_top']} in {100 * g1['tso_forward_frac']:.1f} % of reads; the scaffold anchor follows the 20-nt protospacer after a variable 0-4 nt gap {g1['tso_to_protospacer_gap_top']} whose bases are {gap_desc} (predominantly but not exclusively G; origin - non-templated addition vs. 5' spacer extension - is undocumented); anchors are forward-strand in R1 ({100 * a_frac + 100 * c_frac:.1f} %) and reverse-complement in R2 ({100 * (g2['anchor_A_revcomp_frac'] + g2['anchor_C_revcomp_frac']):.1f} %), so the guide transcript is read 5'->3' in R1. "
      f"Sample indices are in I1/I2 ({fq[fq.read == 'I1'].read_length_mode.iloc[0]} / {fq[fq.read == 'I2'].read_length_mode.iloc[0]} nt). GEX R1 is {x1['read_length_mode']} nt. NOT documented anywhere in the files: the kit / chemistry name for the guide library (directory label 'FB_OTHER_10X'), the capture method (direct capture of the sgRNA via the 5' TSO vs. a separate feature-barcode primer) and whether R2 was intended to be used.",
      g1["file"], "FASTQ parsing (direct)", f"R1 header {g1['header_example']}", "high for positions/orientation (data-derived); low for the intended library chemistry", "official chemistry / kit and read configuration of the guide libraries", True)
    # Q7
    q("Q7", "Guide-counting method: software, version, command, reference, whitelist, correction, mismatch, multi-match, A vs C, source",
      "ANSWERED",
      "Fully documented in code, configuration and logs (see guide_counting_audit.csv): in-house perturbseq_pipeline.guide_counting on guide FASTQ R1; reference = the 560 workbook sequences; cell universe = Cell Ranger filtered GEX barcodes (no whitelist file, no barcode correction); exact UMI de-duplication; exact protospacer match with +/-1 nt shift and an unambiguous 1-mismatch index (current matrices) / exact only (first run); "
      "scaffold class from the 12-nt anchor alternation; Cell Ranger produced no guide counts (GEX-only run, feature_reference null). The pair assignment uses this derived UMI matrix; it is supported by raw-read evidence only in the sense that the matrix is computed from reads - no read-level pairing evidence exists.",
      f"{REPO}/src/perturbseq_pipeline/guide_counting.py; {REPO}/jobs/hanrui_fang/03_count_guides.slurm; {REPO}/results/Hanrui_fang_dual_guide/logs/03_count_guides.log", "source code + job script + run log (direct)",
      "log: 'Guide counting: 8 FASTQ file(s) across 4 sample(s) with 8 worker(s); 560 designed guides, 2 scaffold class(es) [A, C], position_shift=1, max_mismatches=1'", "high", "none about the method; the method's suitability (no barcode correction, filtered-cell universe) is a design choice to review", False)
    # Q8
    q("Q8", "Non-targeting controls: definition of NO-TARGET, NTC-NTC pairs, targeting+NTC validity, documented NTC pairings",
      "PARTIALLY_ANSWERED",
      f"127 workbook records carry the label 'NO-TARGET' in the 'gene' column (the only definition available); read-classified scaffold split {json.dumps(ntc_split)}. Nothing documents whether they were cloned as NTC-NTC constructs, as NTC partners of targeting guides, or both. "
      f"Observed: {mult.get('pair_status_counts', {}).get('pair_non_targeting', 'n/a')} cells with two NTC guides (A+C) and {mult.get('pair_status_counts', {}).get('pair_targeting_plus_ntc', 'n/a')} with one targeting + one NTC guide - the latter class is as large as the same-target class, which is unexpected if every construct were same-target and is unexplained by the files. Treating NTC-NTC cells as the control group and targeting+NTC cells as valid perturbations are both inferences.",
      f"{ctx['xl']} (gene == 'NO-TARGET'); {mult.get('object', '')}", "workbook parse + derived pair statuses", "label 'NO-TARGET' x 127", "high for the label; none for intended use", "intended NTC construct design; whether targeting+NTC constructs exist", True)
    # Q9
    q("Q9", "Sample structure: what do HF011A/B and HF012A/B represent?",
      "PARTIALLY_ANSWERED",
      "Established from files: four separate 10x captures / GEM wells, each with its own Cell Ranger run (sample ids HF011A..HF012B), own GEX library (S29/S31/S33/S35) and own guide library (HF0xxF, S30/S32/S34/S36), sequenced on one flowcell across lanes L004+L005 (lanes are sequencing lanes, not samples). "
      "NOT in any file: whether HF011 vs HF012 are conditions, donors, time points or batches, and whether A/B are biological or technical replicates. The 'condition'/'replicate' labels in config/Hanrui_fang_sample_manifest.csv were written by the analyst from the well ids (condition_status = inferred, replicate_type = unknown).",
      f"{DATA}/HF011A/fastq/SC_5PGEX_10X/...; {REPO}/config/Hanrui_fang_sample_manifest.csv", "FASTQ names / Cell Ranger sample ids (direct) + analyst manifest (computational)", "manifest: condition_status=inferred, replicate_type=unknown, cell_line=unknown", "high for capture structure; none for biological meaning", "meaning of HF011/HF012 and A/B; cell line; treatment", True)
    # Q10
    q("Q10", "Expected guide recovery and MOI: expected vs observed",
      "NOT_ANSWERED",
      "No expected values exist in any file (no protocol, sample sheet, library QC or methods document; Cell Ranger reports are GEX-only). Observed only: 416/560 designed guides detected (144 never observed), scaffold anchor in ~98 % of guide reads, 95.5 % of anchored reads matched to a designed protospacer, 93 % of matched reads carry a called-cell barcode; "
      f"complete-pair (both slots resolved) {100 * mult.get('frac_cells_both_slots_resolved(pair statuses)', float('nan')):.1f} %, > 1 guide {100 * mult.get('frac_cells_gt1_guide_ge3umi', float('nan')):.1f} %, targeting-pair 20.5 %, NTC-pair 15.1 %, targeting+NTC 20.8 % (combined run).",
      f"{REPO}/results/Hanrui_fang_dual_guide/tables/guide_read_retention_funnel.csv; {mult.get('object', '')}", "derived tables (computational)", "guides observed 416/560", "high for observed; none for expected", "expected MOI, expected guide representation, library QC (NGS of the plasmid pool)", True)
    # Q11
    q("Q11", "Experimental modality (KO / CRISPRi / CRISPRa / other)",
      "NOT_ANSWERED",
      f"Undocumented. Keyword search over every text file of the dataset and previous reports: {json.dumps({k: v for k, v in hits_mod.items() if v})[:200]} - the only mentions are in the analyst's own reports stating that the modality is unknown. Cell Ranger used the stock GRCh38-2024-A reference (no Cas9/dCas9/KRAB transgene), so the effector cannot be read out from the expression data either.",
      f"{DATA}/data_audit/reports/10_OPEN_QUESTIONS_AND_RISKS.md; {DATA}/web_summaries/*_web_summary.html (reference)", "text search + Cell Ranger reference (direct)", "no protocol / methods file in the dataset", "high (absence)", "effector (Cas9 / dCas9-KRAB / dCas9-VPR / other) and delivery (lentiviral MOI, selection)", True)
    # Q12
    q("Q12", "Is the current pair-guide model justified?",
      "SUPPORTED_WITH_LIMITATIONS",
      "(1) two scaffold classes: SUPPORTED by direct read evidence; (2) physical dual-guide constructs: INFERRED only (co-occurrence in cells; no construct documentation); (3) exact A-C pairing: NOT SUPPORTED by any file - the same-target rule is an assumption; (4) one pair per cell: NOT SUPPORTED - most cells carry several strong guides per scaffold and only a minority are clean 1A+1C; "
      "(5) NTC pair definition: INFERRED (NO-TARGET label only; targeting+NTC cells unexplained); (6) guide-counting procedure: SUPPORTED (documented, exact matching, unambiguous reference). "
      "Consequently the model can be used for exploratory, clearly-labelled provisional analysis, but a definitive pair-guide biological analysis requires the construct map (or a statement that all same-target combinations are valid), the intended MOI and the NTC design from the experimental team.",
      str(OUT / "answer_matrix.csv"), "synthesis of Q1-Q11", "see decision gate", "n/a", "see questions_for_experimental_team.md", True)
    return pd.DataFrame(A)


# =====================================================================================
# 9. pairing evidence + contradictions
# =====================================================================================
def build_pair_mapping_evidence(ctx):
    rows = []
    xl = ctx["xl"]; wa = ctx["wa"]
    rows.append({"source": str(xl), "evidence_type": "workbook columns", "finding": "no pair / construct / vector / slot / scaffold column", "supports_explicit_pairing": False, "detail": wa[wa.item == "__sheet__"].value.iloc[0][:200]})
    for k, hits in ctx["found"].items():
        rows.append({"source": f"find over {DATA} and {PARENT}", "evidence_type": f"file search: {k}", "finding": f"{len(hits)} file(s)" + (": " + ";".join(hits[:5]) if hits else " - none found"), "supports_explicit_pairing": False, "detail": ""})
    rows.append({"source": f"{DATA}/data_audit + previous reports", "evidence_type": "keyword search over text", "finding": json.dumps(ctx["hits_pair"])[:300], "supports_explicit_pairing": False,
                 "detail": "'pair' mentions are the analyst's own descriptions of the computational pairing rule, not experimental metadata"})
    rows.append({"source": f"{DATA}/web_summaries; Cell Ranger info/outs tars", "evidence_type": "Cell Ranger metadata", "finding": "GEX-only runs (no feature reference, no CRISPR analysis folder); no construct information", "supports_explicit_pairing": False, "detail": ";".join(str(x) for x in ctx["cr"].get("outs_has_crispr_analysis", pd.Series([])).tolist())})
    m = ctx["mult"]
    rows.append({"source": m.get("object", ""), "evidence_type": "cell-level co-occurrence (computational inference)", "finding": f"both scaffold classes in {100 * m.get('frac_cells_ge1A_and_ge1C_ge3umi', 0):.1f} % of cells; exactly 1A+1C in {100 * m.get('frac_cells_exactly_1A_1C_ge3umi', 0):.1f} %; pair statuses {json.dumps(m.get('pair_status_counts', {}))[:250]}",
                 "supports_explicit_pairing": False, "detail": "co-occurrence supports two classes per cell, not which guides are on one construct"})
    ref = REPO / "results/Hanrui_Fang_pair_guide_full/audit/pair_guide_reference.csv"
    if ref.is_file():
        r = pd.read_csv(ref, dtype=str, keep_default_na=False)
        per = r[r.is_non_targeting.str.lower() != "true"].groupby("target_gene_name").agg(nA=("scaffold", lambda s: int((s == "A").sum())), nC=("scaffold", lambda s: int((s == "C").sum())))
        rows.append({"source": str(ref), "evidence_type": "derived reference (previous iteration)", "finding": f"pair_id empty for all 560 guides; per target {per.nA.median():.0f} A and {per.nC.median():.0f} C observed guides -> up to {int((per.nA * per.nC).median())} possible A x C combinations per target (median)", "supports_explicit_pairing": False, "detail": ""})
    return pd.DataFrame(rows)


def build_contradictions(ctx):
    L = []
    sc = ctx["scaffold_df"]; cr = ctx["cr"]; m = ctx["mult"]
    L.append("# Contradictions between files, previous reports and computational assumptions\n")
    L.append("## Contradictions / inconsistencies found\n")
    L.append("1. **Previous reports describe the library as 'dual-guide (scaffold A + C), every vector carries one A and one C guide'.** No dataset file states this; it is an inference from two observed sequence classes and their co-occurrence in cells. The present audit downgrades it to *computational inference* (Q1).")
    L.append("2. **Scaffold class counts differ between reports**: the data_audit subsample stated A 203 / C 199 (or 202), the exact-only full run gave A 202 / C 199 / unknown 159, and the 1-mismatch run gives A 203 / C 201 / unknown 156. The differences come from matching depth (rescued reads classify additional guides), not from the design. None of these numbers is experimentally documented.")
    L.append("3. **'BeginsWithG' has been listed in some earlier tables next to the scaffold class.** The workbook test shows it is exactly `seq[0] == 'G'` (a derived sequence property), so it must not be read as scaffold or cassette information.")
    L.append("4. **The analyst-written sample manifest carries 'condition' and 'replicate' columns** (HF011/HF012, A/B). These are labels derived from well ids (`condition_status = inferred`), not experimental metadata; previous per-well 'reproducibility' statements are well-level only.")
    L.append(f"5. **Same-target pairs vs targeting+NTC pairs**: the pair model assumes same-target constructs, yet cells with one targeting and one NTC guide ({m.get('pair_status_counts', {}).get('pair_targeting_plus_ntc', 'n/a')}) are as frequent as same-target cells ({m.get('pair_status_counts', {}).get('pair_targeting', 'n/a')}). Either targeting+NTC constructs exist by design, or multi-construct cells dominate; the files cannot distinguish these (Q2/Q8).")
    L.append("6. **Guide multiplicity**: the previous data_audit estimated ~48 % 'clean two-guide cells' from a 20 M-read subsample at a fixed 3-UMI threshold, whereas the full-depth matrices show a median of 4 detected guides per cell at 3 UMIs (median 2,762-2,927 guide UMIs per cell). The threshold, not the biology, drives this difference; the audit does not lower any threshold.")
    L.append("7. **Guide 'unknown scaffold' class**: 144-159 designed guides were never observed. Previous reports did not state whether these guides were absent from the cloned library, failed capture, or were excluded upstream; no library QC exists to check.")
    L.append("\n## Gaps (information absent from all files)\n")
    for g in ("construct / vector map linking each scaffold-A guide to a scaffold-C guide", "statement of library architecture (dual cassette vs independent libraries) and promoters", "intended MOI and expected guides per cell",
              "NTC construct design (NTC-NTC, targeting+NTC)", "CRISPR modality / effector and cell line", "meaning of HF011 vs HF012 and A vs B (biological/technical replicate, condition)", "guide-library kit / chemistry and intended read configuration",
              "vendor / plasmid identity and full sequence of the two scaffold variants", "plasmid-pool NGS QC (guide representation, expected recovery)", "any USB-delivered files (none found)"):
        L.append(f"- {g}")
    return "\n".join(L) + "\n"


def build_questions(answers: pd.DataFrame):
    L = ["# Questions for the experimental team\n", "Only questions that the existing files cannot answer are listed (statuses NOT_ANSWERED / PARTIALLY_ANSWERED). Q4 (workbook content) and Q7 (guide-counting method) were answered by the files and are omitted.\n"]
    groups = {
        "Library architecture": ["Is each guide construct a single vector carrying one scaffold-A cassette AND one scaffold-C cassette (dual-guide vector), or are the A- and C-scaffold guides separate libraries / separate vectors that were co-delivered?",
                                 "Which promoters drive the two cassettes, and why were two different scaffold variants used (e.g. to prevent recombination)?"],
        "Guide pairing": ["Please provide the construct map: for every scaffold-A guide, which scaffold-C guide is on the same vector (pair/construct id per oligo). If no map exists, please confirm explicitly whether every same-target A x C combination was cloned (combinatorial) or whether pairs were fixed.",
                          "Are there constructs that combine a targeting guide with a NO-TARGET guide? If so, which ones and what is their intended role?",
                          "For targets with 10 designed guides of which only ~6 are observed, were the unobserved guides excluded from cloning or did they fail?"],
        "Scaffold sequences": ["Please provide the full annotated vector / cassette sequence (both tracr/scaffold variants). The reads show two scaffold classes beginning GTTTAAGAGCTA... and GTTTCAGAGCTA...; please confirm their identity and origin (vendor / plasmid).",
                               "Is a 5'-G prepended (or the first base substituted with G) when the designed protospacer does not start with G? Reads show frequent position-1 G substitutions."],
        "FASTQ read structure": ["Which 10x kit / chemistry and which guide-capture strategy produced the HF0xxF ('FB_OTHER_10X') libraries (5' direct capture of the sgRNA via TSO, or a custom primer)? Please provide the library-construction protocol and the intended read lengths / which reads are informative."],
        "Guide-counting method": ["No question: the counting method is fully documented in the pipeline. (Optional) Do you have Cell Ranger CRISPR or another independent guide quantification to cross-check against?"],
        "Controls": ["How were the 127 NO-TARGET guides intended to be used: as NTC-NTC constructs, as NTC partners of targeting guides, or both? Which NTC-NTC combinations exist?"],
        "Sample / replicate metadata": ["What do HF011 and HF012 represent (condition, donor, time point, batch)? Are A and B biological replicates (independent infections/cultures) or technical replicates (one culture split over two GEM wells)?",
                                        "Which cell line (or donor material) was used, and how were cells selected after transduction?"],
        "Expected MOI and QC": ["What MOI / transduction efficiency was targeted, and how many constructs per cell were expected? Is plasmid-pool NGS QC (guide representation) available for the cloned library?",
                                "Was any guide-level QC (e.g. expected recovery, dropout list) performed before sequencing?"],
        "Experimental modality": ["Which CRISPR modality and effector were used (Cas9 knockout, dCas9-KRAB CRISPRi, dCas9-VPR CRISPRa, other)? Was the effector stably expressed and validated?"],
    }
    for g, qs in groups.items():
        L.append(f"## {g}\n")
        for i, q in enumerate(qs, 1):
            L.append(f"{i}. {q}")
        L.append("")
    L.append("## Minimum information required before a definitive pair-guide analysis\n")
    L.append("1. The A-C construct map (or an explicit statement that all same-target A x C combinations are valid constructs).\n2. Library architecture (dual-cassette vector vs co-delivered libraries) and intended MOI.\n3. NTC construct design.\n4. CRISPR modality/effector and the biological meaning of HF011/HF012 and A/B.")
    return "\n".join(L) + "\n"


# =====================================================================================
# 10. reports
# =====================================================================================
def md_table(df: pd.DataFrame, max_rows=60, max_col=90) -> str:
    if df is None or not len(df):
        return "_none_\n"
    d = df.head(max_rows).copy()
    cols = [str(c) for c in d.columns]
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for _, r in d.iterrows():
        out.append("| " + " | ".join(str(v).replace("|", "\\|").replace("\n", " ")[:max_col] for v in r.tolist()) + " |")
    if len(df) > max_rows:
        out.append(f"\n_{len(df) - max_rows} more rows in the CSV_")
    return "\n".join(out) + "\n"


def html_table(df: pd.DataFrame, max_rows=200) -> str:
    if df is None or not len(df):
        return "<p><i>none</i></p>"
    return df.head(max_rows).to_html(index=False, escape=True, border=0, classes="t", na_rep="")


def write_reports(ctx, answers, pme, contradictions_md, questions_md, provenance):
    sc = ctx["scaffold_df"]; fq = ctx["fq"]; wa = ctx["wa"]; cr = ctx["cr"]; gm = ctx["gm"]; cm = ctx["cm"]; sm = ctx["sm"]; inv = ctx["inv"]; usb = ctx["usb"]; m = ctx["mult"]
    status_counts = answers["status"].value_counts().to_dict()
    executive = [
        "**Executive conclusion.** The existing files establish (i) two scaffold-associated sequence classes in the guide reads, (ii) a fully documented in-house guide-counting procedure, (iii) the exact content of the design workbook, and (iv) that the four samples are separate 10x captures. "
        "They do **not** establish that scaffold-A and scaffold-C guides are physically paired on one construct, which A guide pairs with which C guide, that one pair per cell was intended, how NO-TARGET guides were meant to be used, the CRISPR modality, or the biological meaning of HF011/HF012 and A/B. "
        f"The current pair-guide model is classified **{answers.loc[answers.question_id == 'Q12', 'status'].iloc[0]}**. Status counts: {status_counts}. "
        "**Decision gate: the current files are NOT sufficient for a definitive pair-guide biological analysis.** Exploratory, clearly labelled provisional analysis may continue; final perturbation claims must wait for the construct map, the intended MOI and the NTC design from the experimental team.",
    ]
    established = ["Two scaffold-associated sequence classes exist in the guide reads (anchors GTTTAAGAGCTA / GTTTCAGAGCTA; read-derived consensus of the sequenced ~66 nt of each scaffold in scaffold_evidence.csv; the insert ends in the Illumina adapter).",
                   "Guide-library R1 read layout (16-nt cell barcode = Cell Ranger barcode, 12-nt UMI, TSO, non-templated G run, 20-nt protospacer, scaffold) - derived directly from reads.",
                   "Workbook content: 560 unique 20-nt sequences, 46 labels, 127 NO-TARGET, no scaffold/pair/construct/sample columns; BeginsWithG = seq[0]=='G'.",
                   "Guide counting method, reference, cell universe, UMI handling, mismatch rules, scaffold classification (code + logs).",
                   "Cell Ranger runs were GEX-only (no feature reference, no CRISPR output); guide counts exist only as pipeline-derived matrices.",
                   "Four separate GEM wells / captures, each with its own GEX and guide library, sequenced on one flowcell (lanes L004/L005).",
                   "No USB-derived folder or file exists under the data parent directory or the repository."]
    inferred = ["'Scaffold A' / 'scaffold C' labels and the per-guide scaffold class (majority anchor of matched reads).", "Dual-cassette (one A + one C guide per construct) architecture.", "Same-target A+C combinations as valid pairs; NTC-NTC cells as controls; targeting+NTC cells as not assignable.",
                "Condition (HF011/HF012) and replicate (A/B) labels in the sample manifest.", "Guide-detection threshold (3 UMIs) and dominance ratio (2.0) as the definition of a 'strong' guide."]
    unknown = ["Construct / pair map; library architecture and promoters.", "Intended MOI; expected guides per cell; plasmid-pool QC.", "NTC construct design.", "CRISPR modality / effector; cell line; selection.", "Meaning of HF011 vs HF012 and A vs B.",
               "Guide-library kit / chemistry and intended read configuration.", "Vendor / plasmid identity of the two scaffold variants."]
    # ---------- Markdown
    M = [f"# Hanrui Fang guide-library and pairing audit\n", f"SLURM job {job_id}; {provenance['timestamp']}; branch `{provenance['git_branch']}` commit `{provenance['git_commit']}`; conda env `{provenance['conda_env']}`; host {provenance['host']}.\n",
         "## 1. Executive conclusion\n", *executive, "",
         "## 2. Files inspected\n", f"{len(inv)} files inventoried (`file_inventory.csv`); categories: {json.dumps(ctx['cats'])}. Searches for FASTA / feature reference / whitelist / sample sheet / protocol files: " + "; ".join(f"{k}: {len(v)}" for k, v in ctx["found"].items()) + ".\n",
         "## 3. USB-folder discovery result\n", usb["statement"] + f" (searched: {usb['search_roots']})\n",
         "## 4. Workbook structure and interpretation\n", md_table(wa, 60, 160),
         "## 5. FASTA/FASTQ scaffold evidence\n", "No FASTA files exist in the dataset. Read-derived scaffold consensus per library and class:\n", md_table(sc[["well", "scaffold_class", "anchor", "n_reads_with_anchor", "observed_scaffold_consensus_from_anchor", "observed_scaffold_length_nt", "insert_ends_in_adapter", "mean_per_position_agreement_first_60", "closest_reference_scaffold(annotation_only)", "identity_to_closest", "match_call"]], 8, 110),
         "\nRead-structure statistics (200k reads per file; `fastq_read_structure_audit.csv`):\n", md_table(fq[["library_type", "sample", "read", "read_length_mode", "frac_first16_in_gex_filtered_barcodes", "tso_forward_frac", "tso_position_top", "anchor_A_forward_frac", "anchor_C_forward_frac", "anchor_canonical_T_forward_frac", "anchor_A_revcomp_frac", "anchor_C_revcomp_frac", "tso_to_protospacer_gap_top"]].fillna(""), 40, 60),
         "## 6. Guide-counting evidence\n", md_table(gm, 20, 200), "\nDerived count matrices:\n", md_table(cm[["well", "n_features", "n_barcodes", "barcodes_match_gex_filtered", "reads_total", "reads_with_scaffold_anchor", "reads_spacer_matched", "reads_via_1mismatch", "unique_cell_guide_umis", "reads_scaffold_A", "reads_scaffold_C"]], 4, 60),
         f"\nObserved multiplicity (combined QC-pass object): `{json.dumps({k: v for k, v in m.items() if k not in ('pair_status_counts', 'object')}, default=str)}`; pair statuses: `{json.dumps(m.get('pair_status_counts', {}))}`.\n",
         "## 7. Pairing evidence\n", md_table(pme, 20, 200),
         "## 8. Sample and replicate evidence\n", md_table(sm[["sample", "guide_library", "flowcell_run", "sequencing_lanes", "gex_sample_index_number", "guide_sample_index_number", "cellranger_chemistry", "n_cells_cellranger", "manifest_condition", "manifest_replicate", "manifest_replicate_type", "manifest_condition_status", "classification"]], 4, 160),
         "\nCell Ranger metadata per well:\n", md_table(cr[[c for c in cr.columns if c.startswith(("well", "tbl_", "json_", "outs_has", "info_tar_members", "mentions"))]].fillna(""), 4, 80),
         "## 9. Question-by-question answer matrix\n", md_table(answers, 12, 700),
         "## 10. Contradictions between files and previous reports\n", contradictions_md,
         "## 11. What is established\n", "\n".join(f"- {x}" for x in established), "\n## 12. What is only inferred\n", "\n".join(f"- {x}" for x in inferred), "\n## 13. What remains unknown\n", "\n".join(f"- {x}" for x in unknown),
         "\n## 14. Exact questions to send to the experimental team\n", questions_md,
         "## 15. Decision gate\n", "**Can a definitive pair-guide analysis proceed on the current files? No.** The two-class scaffold evidence and the counting procedure are solid, but exact pairing, one-pair-per-cell intent, NTC design and modality are undocumented. "
         "Provisional, clearly-labelled analyses (as produced in results/Hanrui_Fang_pair_guide_full) remain exploratory. No thresholds were changed in this audit and no perturbation claims are made.\n"]
    (OUT / "library_design_audit.md").write_text("\n".join(M))
    # ---------- HTML
    css = "body{font-family:system-ui,sans-serif;margin:24px;max-width:1500px;color:#222}table.t{border-collapse:collapse;font-size:11px}table.t td,table.t th{border:1px solid #ddd;padding:2px 6px;vertical-align:top}h2{border-bottom:1px solid #ccc;margin-top:32px}.gate{background:#fff4e5;border-left:5px solid #eda100;padding:10px}.ok{color:#1baf7a}.bad{color:#e34948}nav a{margin-right:10px}pre{background:#f6f6f6;padding:8px;white-space:pre-wrap;font-size:11px}"
    secs = ["Executive conclusion", "Files inspected", "USB-folder discovery result", "Workbook structure and interpretation", "FASTA/FASTQ scaffold evidence", "Guide-counting evidence", "Pairing evidence", "Sample and replicate evidence", "Question-by-question answer matrix",
            "Contradictions between files and previous reports", "What is established", "What is only inferred", "What remains unknown", "Exact questions to send to the experimental team", "Decision gate"]
    def md2html(t):
        t = html.escape(t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t); t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
        t = re.sub(r"(?m)^### (.+)$", r"<h4>\1</h4>", t); t = re.sub(r"(?m)^## (.+)$", r"<h3>\1</h3>", t); t = re.sub(r"(?m)^# (.+)$", r"<h3>\1</h3>", t)
        t = re.sub(r"(?m)^- (.+)$", r"<li>\1</li>", t); t = re.sub(r"(?m)^(\d+)\. (.+)$", r"<li>\2</li>", t)
        return "<p>" + t.replace("\n\n", "</p><p>").replace("\n", "<br>") + "</p>"
    H = [f"<html><head><meta charset='utf-8'><title>Hanrui Fang guide-library and pairing audit</title><style>{css}</style></head><body>",
         f"<h1>Hanrui Fang guide-library and pairing audit</h1><p>SLURM job <b>{job_id}</b> · {provenance['timestamp']} · branch <code>{provenance['git_branch']}</code> commit <code>{provenance['git_commit']}</code> · conda env <code>{provenance['conda_env']}</code> · host {provenance['host']} · evidence report only (no perturbation analysis, no threshold changes).</p>",
         "<nav>" + "".join(f"<a href='#s{i}'>{i}. {s}</a>" for i, s in enumerate(secs, 1)) + "</nav>"]
    H.append(f"<h2 id='s1'>1. {secs[0]}</h2><div class='gate'>{md2html(executive[0])}</div>")
    H.append(f"<h2 id='s2'>2. {secs[1]}</h2><p>{len(inv)} files inventoried; categories: <code>{html.escape(json.dumps(ctx['cats']))}</code>. File searches: " + "; ".join(f"{k}: {len(v)}" for k, v in ctx["found"].items()) + ".</p>" + html_table(inv[inv.category != "previous_audit_other"], 400))
    H.append(f"<h2 id='s3'>3. {secs[2]}</h2><p><b>{html.escape(usb['statement'])}</b><br>Search roots: {html.escape(str(usb['search_roots']))}</p>")
    H.append(f"<h2 id='s4'>4. {secs[3]}</h2>" + html_table(wa, 100))
    H.append(f"<h2 id='s5'>5. {secs[4]}</h2><p>No FASTA / feature-reference files exist in the dataset; scaffold sequences below are <b>read-derived consensuses</b> (direct evidence for the sequence classes) compared with published scaffolds (annotation only). The 12-nt anchors in the pipeline configuration are computational choices.</p>" + html_table(sc) + "<h4>Read-structure statistics (200k reads per file)</h4>" + html_table(fq.fillna("")))
    H.append(f"<h2 id='s6'>6. {secs[5]}</h2>" + html_table(gm) + "<h4>Derived count matrices</h4>" + html_table(cm) + f"<h4>Observed guide multiplicity (combined QC-pass object)</h4><pre>{html.escape(json.dumps(m, indent=1, default=str))}</pre><h4>Counter docstring (source of truth for the read model)</h4><pre>{html.escape(ctx['doc'][:3000])}</pre><h4>Counter configuration used (config/hanrui_fang.yaml)</h4><pre>{html.escape(ctx['fq_block'])}</pre>")
    H.append(f"<h2 id='s7'>7. {secs[6]}</h2>" + html_table(pme))
    H.append(f"<h2 id='s8'>8. {secs[7]}</h2>" + html_table(sm) + "<h4>Cell Ranger metadata</h4>" + html_table(cr.fillna("")))
    H.append(f"<h2 id='s9'>9. {secs[8]}</h2><p>Status counts: <code>{html.escape(json.dumps(status_counts))}</code></p>" + html_table(answers))
    H.append(f"<h2 id='s10'>10. {secs[9]}</h2>" + md2html(contradictions_md))
    H.append(f"<h2 id='s11'>11. {secs[10]}</h2><ul>" + "".join(f"<li class='ok'>{html.escape(x)}</li>" for x in established) + "</ul>")
    H.append(f"<h2 id='s12'>12. {secs[11]}</h2><ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in inferred) + "</ul>")
    H.append(f"<h2 id='s13'>13. {secs[12]}</h2><ul>" + "".join(f"<li class='bad'>{html.escape(x)}</li>" for x in unknown) + "</ul>")
    H.append(f"<h2 id='s14'>14. {secs[13]}</h2>" + md2html(questions_md))
    H.append(f"<h2 id='s15'>15. {secs[14]}</h2><div class='gate'><b>Are the current files sufficient for a definitive pair-guide analysis? NO.</b><br>" + md2html(M[-1].split("**", 2)[-1]) + "</div>")
    H.append("<h3>Evidence extracts</h3>" + html_table(pd.DataFrame(evidence_extracts), 300) + "<h3>Provenance</h3><pre>" + html.escape(json.dumps(provenance, indent=1)) + "</pre></body></html>")
    (OUT / "library_design_audit.html").write_text("\n".join(H))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    t0 = datetime.now()
    usb, found, cats = do_inventory()
    texts, hits_mod, hits_pair, claims = do_previous_reports()
    # keyword hits restricted to dataset (non-audit) text files
    ds_text = "\n".join(p.read_text(errors="ignore") for p in DATA.rglob("*") if p.is_file() and "data_audit" not in p.parts and p.suffix in (".txt", ".md5", ".csv", ".json", ".md"))
    hits_pair_files = {kw: len(re.findall(r"\b" + re.escape(kw) + r"\b", ds_text, flags=re.IGNORECASE)) for kw in KEYWORDS_PAIR}
    cr = do_cellranger()
    xl, wa, guides, seq, labels, ntc_mask = do_workbook()
    gex_barcodes = {}
    for w in WELLS:
        with gzip.open(REPO / "results/outputs/hanrui_fang_inputs/cellranger" / w / "filtered_feature_bc_matrix/barcodes.tsv.gz", "rt") as fh:
            gex_barcodes[w] = {l.strip().split("-")[0] for l in fh if l.strip()}
    fq, scaffold_df = do_fastq(gex_barcodes)
    cm, mult, gm, stats_all, doc, fq_block, log03 = do_counts_and_method(gex_barcodes)
    sm, mdf = do_samples(cr, fq)
    inv = pd.DataFrame(inventory).drop_duplicates("path")
    ctx = dict(usb=usb, found=found, cats=dict(cats), cr=cr, xl=xl, wa=wa, fq=fq, scaffold_df=scaffold_df, cm=cm, mult=mult, gm=gm, sm=sm, inv=inv, hits_mod=hits_mod, hits_pair=hits_pair, hits_pair_files=hits_pair_files, doc=doc, fq_block=fq_block)
    answers = build_answers(ctx)
    pme = build_pair_mapping_evidence(ctx)
    contradictions_md = build_contradictions(ctx)
    questions_md = build_questions(answers)
    git = {k: sh(cmd, cwd=REPO).strip() for k, cmd in (("git_branch", "git rev-parse --abbrev-ref HEAD"), ("git_commit", "git rev-parse HEAD"), ("git_status", "git status --short"))}
    provenance = {"timestamp": t0.strftime("%Y-%m-%d %H:%M:%S"), "slurm_job_id": job_id, "slurm_job_name": os.environ.get("SLURM_JOB_NAME"), "node": os.environ.get("SLURMD_NODENAME"), "host": os.uname().nodename,
                  "cpus": os.environ.get("SLURM_CPUS_PER_TASK"), "mem": os.environ.get("SLURM_MEM_PER_NODE"), "conda_env": os.environ.get("CONDA_DEFAULT_ENV"), "python": sys.version.split()[0], "conda_prefix": os.environ.get("CONDA_PREFIX"),
                  **git, "dataset_root": str(DATA), "workbook": str(xl), "workbook_md5": hashlib.md5(xl.read_bytes()).hexdigest(), "n_files_inventoried": int(len(inv)), "fastq_reads_sampled_per_file": N_SAMPLE_READS, "fastq_reads_for_scaffold_consensus": N_CONSENSUS_READS,
                  "usb_search": usb, "file_searches": {k: len(v) for k, v in found.items()}, "previous_reports_read": [p for p in texts], "thresholds_changed": False, "perturbation_analysis_run": False, "previous_results_overwritten": False,
                  "elapsed_seconds": None}
    # write CSVs
    inv.to_csv(OUT / "file_inventory.csv", index=False)
    wa.to_csv(OUT / "workbook_audit.csv", index=False)
    scaffold_df.to_csv(OUT / "scaffold_evidence.csv", index=False)
    pme.to_csv(OUT / "pair_mapping_evidence.csv", index=False)
    fq.to_csv(OUT / "fastq_read_structure_audit.csv", index=False)
    gm.to_csv(OUT / "guide_counting_audit.csv", index=False)
    cm.to_csv(OUT / "guide_count_matrix_audit.csv", index=False)
    sm.to_csv(OUT / "sample_metadata_audit.csv", index=False)
    cr.to_csv(OUT / "cellranger_metadata_audit.csv", index=False)
    answers.to_csv(OUT / "answer_matrix.csv", index=False)
    pd.DataFrame(evidence_extracts).to_csv(OUT / "evidence_extracts.csv", index=False)
    pd.DataFrame(claims).to_csv(OUT / "previous_report_claims_checked.csv", index=False)
    json.dump(mult, open(OUT / "observed_multiplicity.json", "w"), indent=1, default=str)
    (OUT / "contradictions_and_gaps.md").write_text(contradictions_md)
    (OUT / "questions_for_experimental_team.md").write_text(questions_md)
    provenance["elapsed_seconds"] = (datetime.now() - t0).total_seconds()
    json.dump(provenance, open(OUT / "provenance.json", "w"), indent=1, default=str)
    write_reports(ctx, answers, pme, contradictions_md, questions_md, provenance)
    readme = ["# Hanrui Fang guide-library and pairing audit", "", f"Evidence audit (SLURM job {job_id}, {provenance['timestamp']}) of the guide design before any final pair-guide analysis. No thresholds changed, no perturbation analysis run, no previous results touched.", "",
              "- `library_design_audit.md` / `library_design_audit.html`: full report (15 sections; HTML self-contained)", "- `answer_matrix.csv`: Q1-Q12 with status, evidence path/type, exact item, confidence, uncertainty, confirmation flag",
              "- `file_inventory.csv`, `workbook_audit.csv`, `scaffold_evidence.csv`, `pair_mapping_evidence.csv`, `fastq_read_structure_audit.csv`, `guide_counting_audit.csv`, `guide_count_matrix_audit.csv`, `sample_metadata_audit.csv`, `cellranger_metadata_audit.csv`, `evidence_extracts.csv`, `previous_report_claims_checked.csv`, `observed_multiplicity.json`",
              "- `contradictions_and_gaps.md`, `questions_for_experimental_team.md`, `provenance.json`, `slurm/`", "", f"Statuses: {answers['status'].value_counts().to_dict()}; Q12 pair model: {answers.loc[answers.question_id == 'Q12', 'status'].iloc[0]}."]
    (OUT / "README.md").write_text("\n".join(readme) + "\n")
    log(answers[["question_id", "status"]].to_string(index=False))
    log("done in", provenance["elapsed_seconds"], "s")


if __name__ == "__main__":
    main()
