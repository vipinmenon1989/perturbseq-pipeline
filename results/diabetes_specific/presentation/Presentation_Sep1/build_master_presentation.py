# -*- coding: utf-8 -*-
"""
Master Scientific Presentation Builder
Generates: PertTF_Master_Scientific_Story.pptx
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# -----------------------------------------------------------------------------
# Color Palette & Typography
# -----------------------------------------------------------------------------
NAVY_PRIMARY   = RGBColor(26, 43, 76)     # #1A2B4C - Deep Navy for main titles
SLATE_ACCENT   = RGBColor(13, 110, 253)   # #0D6EFD - Vibrant Blue for tags / highlights
CYAN_DARK      = RGBColor(0, 102, 153)    # #006699 - Deep Cyan for section labels
CHARCOAL_BODY  = RGBColor(45, 55, 72)     # #2D3748 - Dark Charcoal for readable body text
MUTED_GRAY     = RGBColor(100, 116, 139)  # #64748B - Slate Gray for subtitles / metadata
CARD_BG        = RGBColor(248, 249, 250)  # #F8F9FA - Light Gray card background
CARD_BORDER    = RGBColor(226, 232, 240)  # #E2E8F0 - Subtle slate border
TAKEAWAY_BG    = RGBColor(239, 246, 255)  # #EFF6FF - Soft blue takeaway container
TAKEAWAY_TEXT  = RGBColor(26, 54, 93)     # #1A365D - Deep navy text for takeaways
WHITE          = RGBColor(255, 255, 255)

FONT_HEADING   = "Arial"
FONT_BODY      = "Calibri"

# Slide Dimensions (16:9 Widescreen)
SLIDE_WIDTH_IN  = 13.333
SLIDE_HEIGHT_IN = 7.5

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = prs.slide_layouts[6] # Blank layout
    
    # -------------------------------------------------------------------------
    # Helper Functions
    # -------------------------------------------------------------------------
    def add_header(slide, tag_text, title_text, subtitle_text=None):
        # Section Tag Box
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = FONT_HEADING
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = CYAN_DARK
        
        # Action Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY_PRIMARY
        
        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.3))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = MUTED_GRAY

    def add_fitted_image(slide, image_path, left_in, top_in, max_w_in, max_h_in, add_border=True):
        if not os.path.exists(image_path):
            print(f"[WARN] Image not found: {image_path}")
            return None
        
        # Background card for figure
        if add_border:
            card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_in), Inches(top_in), Inches(max_w_in), Inches(max_h_in))
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            card.line.color.rgb = CARD_BORDER
            card.line.width = Pt(1)
        
        # Get dimensions and compute contain fit
        im = Image.open(image_path)
        im_w, im_h = im.size
        scale = min((max_w_in - 0.2) / im_w, (max_h_in - 0.2) / im_h)
        fit_w = im_w * scale
        fit_h = im_h * scale
        
        pos_left = left_in + (max_w_in - fit_w) / 2.0
        pos_top  = top_in  + (max_h_in - fit_h) / 2.0
        
        pic = slide.shapes.add_picture(image_path, Inches(pos_left), Inches(pos_top), width=Inches(fit_w), height=Inches(fit_h))
        return pic

    def add_interpretation_card(slide, left_in, top_in, width_in, height_in, question_text, bullets, takeaway_text=None):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)
        
        # Text Frame
        tb = slide.shapes.add_textbox(Inches(left_in + 0.25), Inches(top_in + 0.2), Inches(width_in - 0.5), Inches(height_in - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Question Header
        p_q = tf.paragraphs[0]
        p_q.text = question_text
        p_q.font.name = FONT_HEADING
        p_q.font.size = Pt(13)
        p_q.font.bold = True
        p_q.font.color.rgb = CYAN_DARK
        p_q.space_after = Pt(10)
        
        # Bullets
        for b_title, b_desc in bullets:
            p_b = tf.add_paragraph()
            p_b.font.name = FONT_BODY
            p_b.font.size = Pt(12)
            p_b.space_after = Pt(8)
            
            run_title = p_b.add_run()
            run_title.text = f"• {b_title}: "
            run_title.font.bold = True
            run_title.font.color.rgb = NAVY_PRIMARY
            
            run_desc = p_b.add_run()
            run_desc.text = b_desc
            run_desc.font.color.rgb = CHARCOAL_BODY
        
        # Takeaway Box
        if takeaway_text:
            p_t = tf.add_paragraph()
            p_t.space_before = Pt(8)
            p_t.font.name = FONT_BODY
            p_t.font.size = Pt(11.5)
            
            run_lbl = p_t.add_run()
            run_lbl.text = "Takeaway: "
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = SLATE_ACCENT
            
            run_txt = p_t.add_run()
            run_txt.text = takeaway_text
            run_txt.font.italic = True
            run_txt.font.color.rgb = NAVY_PRIMARY

    print("Building slides...")

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_WIDTH_IN), Inches(SLIDE_HEIGHT_IN))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_PRIMARY
    bg1.line.fill.background()
    
    # Title Text
    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p1_tag = tf1.paragraphs[0]
    p1_tag.text = "COMPUTATIONAL BIOLOGY & FUNCTIONAL GENOMICS"
    p1_tag.font.name = FONT_HEADING
    p1_tag.font.size = Pt(13)
    p1_tag.font.bold = True
    p1_tag.font.color.rgb = SLATE_ACCENT
    p1_tag.space_after = Pt(14)
    
    p1_title = tf1.add_paragraph()
    p1_title.text = "Context-Aware AI Modeling for Genome-Scale Perturbation Prediction Across Pancreatic Lineages and Disease States"
    p1_title.font.name = FONT_HEADING
    p1_title.font.size = Pt(28)
    p1_title.font.bold = True
    p1_title.font.color.rgb = WHITE
    p1_title.space_after = Pt(18)
    
    p1_sub = tf1.add_paragraph()
    p1_sub.text = "From Single-Cell Knockout Villages to pertTF and Multi-Dimensional Phenotypic Decomposition"
    p1_sub.font.name = FONT_BODY
    p1_sub.font.size = Pt(16)
    p1_sub.font.color.rgb = RGBColor(203, 213, 225)
    p1_sub.space_after = Pt(24)
    
    p1_auth = tf1.add_paragraph()
    p1_auth.text = "Yangqi Su*, Dingyu Liu*, Vipin Menon*, Danwei Huangfu#, Wei Li# | CHOP, UPenn, MSKCC, Weill Cornell"
    p1_auth.font.name = FONT_BODY
    p1_auth.font.size = Pt(13)
    p1_auth.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: Scientific Overview / Conceptual Premise
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "EXECUTIVE FRAMING", "Traversing Causal Developmental Genetics, Deep Learning, and Phenotyping", "Three integrated pillars bridging biological discovery, machine learning prediction, and multi-scale phenotyping")
    
    # 3 Strategy Columns
    col_w = 3.65
    col_gap = 0.38
    col_top = 1.8
    col_h = 4.9
    
    # Col 1: Biology
    c1 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(col_top), Inches(col_w), Inches(col_h))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tb_c1 = s2.shapes.add_textbox(Inches(1.0), Inches(col_top + 0.25), Inches(col_w - 0.4), Inches(col_h - 0.5))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "1. BIOLOGICAL FOUNDATION"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = CYAN_DARK
    p.space_after = Pt(10)
    p = tf_c1.add_paragraph()
    p.text = "• Longitudinal Knockout Village: 79 hPSC lines targeting 30 disease genes across 5 differentiation stages (>111,000 single cells, 14 cell types)."
    p.font.size = Pt(12)
    p.space_after = Pt(8)
    p = tf_c1.add_paragraph()
    p.text = "• Lineage Rewiring: Genetic perturbations causally divert endocrine progenitors into alternative enterochromaffin (SC-EC) fates at the expense of SC-beta."
    p.font.size = Pt(12)
    p.space_after = Pt(8)
    p = tf_c1.add_paragraph()
    p.text = "• The Combinatorial Barrier: Testing ~20,000 genes across stages is experimentally intractable, requiring transferable AI models."
    p.font.size = Pt(12)

    # Col 2: pertTF
    c2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + col_w + col_gap), Inches(col_top), Inches(col_w), Inches(col_h))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = SLATE_ACCENT
    c2.line.width = Pt(1.5)
    tb_c2 = s2.shapes.add_textbox(Inches(1.0 + col_w + col_gap), Inches(col_top + 0.25), Inches(col_w - 0.4), Inches(col_h - 0.5))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "2. pertTF AI MODEL"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = SLATE_ACCENT
    p.space_after = Pt(10)
    p = tf_c2.add_paragraph()
    p.text = "• Context-Aware Architecture: Transformer encoder with NB-NLL loss, supervised contrastive learning, and gene graph priors."
    p.font.size = Pt(12)
    p.space_after = Pt(8)
    p = tf_c2.add_paragraph()
    p.text = "• Higher-Order Phenotypes: Direct regression of lochNESS scores captures population shifts beyond point-wise gene expression."
    p.font.size = Pt(12)
    p.space_after = Pt(8)
    p = tf_c2.add_paragraph()
    p.text = "• Generalization & Translation: Outperforms foundation models in unseen cell contexts, generalizes to unseen genes, transfers to clinical T2D islets, and enables virtual genetic screens."
    p.font.size = Pt(12)

    # Col 3: WIP
    c3 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + (col_w + col_gap)*2), Inches(col_top), Inches(col_w), Inches(col_h))
    c3.fill.solid()
    c3.fill.fore_color.rgb = CARD_BG
    c3.line.color.rgb = CARD_BORDER
    tb_c3 = s2.shapes.add_textbox(Inches(1.0 + (col_w + col_gap)*2), Inches(col_top + 0.25), Inches(col_w - 0.4), Inches(col_h - 0.5))
    tf_c3 = tb_c3.text_frame
    tf_c3.word_wrap = True
    p = tf_c3.paragraphs[0]
    p.text = "3. MULTI-SCALE PHENOTYPING"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = CYAN_DARK
    p.space_after = Pt(10)
    p = tf_c3.add_paragraph()
    p.text = "• Decompressing End-to-End Prediction: Disentangles single-cell response penetrance (PS), multivariate magnitude (Energy Distance), and state localization (signed lochNESS)."
    p.font.size = Pt(12)
    p.space_after = Pt(8)
    p = tf_c3.add_paragraph()
    p.text = "• Decoupling Discovered: PS couples strongly with Energy Distance (rho = +0.641), but decouples from directional state localization."
    p.font.size = Pt(12)
    p.space_after = Pt(8)
    p = tf_c3.add_paragraph()
    p.text = "• Future Foundation Models: Guides the next generation of phenotype-aware perturbation architectures."
    p.font.size = Pt(12)

    # =========================================================================
    # SECTION 1: BIOLOGICAL FOUNDATION (Slides 3-9)
    # =========================================================================
    
    # SLIDE 3: Human Pancreatic Differentiation Models Causal Disease Checkpoints
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "BIOLOGICAL FOUNDATION", "Human Pancreatic Differentiation Models Causal Disease Checkpoints", "Recapitulating embryonic islet development to study monogenic diabetes and lineage regulation")
    add_fitted_image(s3, "clean_sources/precursor_figures_png/Fig.1-1.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s3, 8.6, 1.8, 4.0, 5.0,
        "Why Pancreatic Differentiation?",
        [
            ("Developmental Checkpoints", "Step-wise in vitro differentiation (Day 0 ESC -> Day 3 DE -> Day 7 PFG -> Day 11 EnP -> Day 18 SC-islet) faithfully mirrors human fetal islet specification."),
            ("Disease Gene Convergence", "Monogenic diabetes genes (MODY, neonatal diabetes) and T2D GWAS loci encode essential transcription factors governing these transitions."),
            ("Experimental Objective", "Causally dissect how loss-of-function mutations alter developmental trajectories at single-cell resolution.")
        ],
        "hPSC differentiation provides an isogenic human platform to study stage-specific developmental genetic lesions."
    )

    # SLIDE 4: Single-Cell Knockout Village Design
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "BIOLOGICAL FOUNDATION", "A Single-Cell Knockout Village Profiles 30 Regulators Across 5 Stages", "Pooled competition culture eliminates batch effects and captures clonal dynamics across differentiation")
    add_fitted_image(s4, "clean_sources/precursor_figures_png/Fig.1-1.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s4, 8.6, 1.8, 4.0, 5.0,
        "Knockout Village Framework",
        [
            ("Library Scale", "79 uniquely barcoded, genotype-verified clonal hPSC lines targeting 30 transcription factors and chromatin remodelers."),
            ("Competitive Co-Culture", "Pooled mutant village was co-cultured with 70% unlabeled WT cells to buffer against non-cell-autonomous paracrine artifacts."),
            ("Single-Cell Resolution", "Profiled 111,581 single cells across 5 stages, resolving 14 distinct cell lineages and high clone-to-clone concordance.")
        ],
        "The knockout village generates a robust, longitudinal dataset linking genetic null mutations to single-cell lineage choices."
    )

    # SLIDE 5: Loss of Lineage Regulators Impairs Beta-Cell Formation and State
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "BIOLOGICAL FOUNDATION", "Loss of Lineage Regulators Impairs Beta-Cell Formation and Maturation", "Mutations trigger either complete developmental blockage or functional maturation failure in residual cells")
    add_fitted_image(s5, "clean_sources/precursor_figures_png/Fig.2-1.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s5, 8.6, 1.8, 4.0, 5.0,
        "Two Distinct Pathological Modes",
        [
            ("Complete Lineage Collapse", "PDX1, RFX6, PAX6, NEUROD1, and GLIS3 knockouts cause near-total elimination of SC-beta cells (<1-2% of WT yield)."),
            ("Maturation & Identity Loss", "Residual beta cells in HNF4A and NEUROD1 mutants display severe downregulation of INS, MAFA, SLC30A8, and G6PC2, accompanied by ectopic progenitor markers."),
            ("Critical Observation", "Total endocrine commitment (CHGA+) is preserved in several mutants, indicating cells did not simply undergo apoptosis.")
        ],
        "Loss of beta-cell fate is accompanied by surviving mutant endocrine cells, raising the question of where these cells diverted."
    )

    # SLIDE 6: Lineage Diversion into Non-Pancreatic Fates
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "BIOLOGICAL FOUNDATION", "Perturbations Divert Progenitors into Alternative Non-Pancreatic Lineages", "Early transcription factor knockouts redirect endodermal progenitors into endothelial and hepatic fates")
    add_fitted_image(s6, "clean_sources/precursor_figures_png/Fig.3-1.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s6, 8.6, 1.8, 4.0, 5.0,
        "Active Lineage Transdifferentiation",
        [
            ("GATA6 -> Endothelial", "GATA6-/- cells completely fail definitive endoderm commitment and actively transdifferentiate into endothelial-like cells (CD34+, PECAM1+)."),
            ("FOXA2 -> Hepatic Progenitors", "FOXA2-/- cells stall in posterior foregut patterning and divert into liver/hepatic lineages expressing ALB, AFP, and APOA1."),
            ("HHEX -> Gut Tube Arrest", "HHEX-/- and enhancer knockouts arrest at early primitive gut tube (PGT) stages.")
        ],
        "Developmental failure in early TF mutants involves active diversion into non-pancreatic lineages rather than passive death."
    )

    # SLIDE 7: Reciprocal SC-Beta vs SC-EC Fate Competition
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "BIOLOGICAL FOUNDATION", "RFX6, PDX1, and PAX6 Regulate a Reciprocal SC-Beta vs SC-EC Fate Choice", "Late endocrine knockouts stoichiometrically expand serotonergic enterochromaffin-like cells at the expense of beta cells")
    add_fitted_image(s7, "clean_sources/precursor_figures_png/Fig.4-1.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s7, 8.6, 1.8, 4.0, 5.0,
        "The SC-Beta vs SC-EC Trade-Off",
        [
            ("Stoichiometric Fate Switch", "In WT, SC-beta represents ~45% and SC-EC represents ~15% of endocrine cells. In RFX6-/-, PDX1-/-, and PAX6-/-, SC-beta collapses to <2% while SC-EC expands to >70-80%."),
            ("Serotonergic Identity", "SC-EC cells express serotonin biosynthesis machinery (SLC18A1, TPH1, DDC) and neuronal projection programs (Program h3)."),
            ("Gatekeeper Role", "RFX6, PDX1, and PAX6 act as developmental gatekeepers that actively repress SC-EC fate to safeguard beta-cell specification.")
        ],
        "SC-beta and SC-EC represent mutually exclusive, competing endocrine fates branching from NEUROG3+ endocrine progenitors."
    )

    # SLIDE 8: Predictive Modeling Identifies ISL1 as a Key Repressor
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "BIOLOGICAL FOUNDATION", "Predictive Modeling Identifies ISL1 as a Key Repressor Safeguarding Beta Fate", "Integrating mutant response ratios predicts causal regulators, validated by knockout and overexpression rescue")
    add_fitted_image(s8, "clean_sources/precursor_pdf_pages/page-44.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s8, 8.6, 1.8, 4.0, 5.0,
        "Predictive Modeling & ISL1 Discovery",
        [
            ("Linear Regression Discovery", "Regression of EnP transcription factors against SC-EC/SC-beta ratios prioritized ISL1 as the top negative predictor (p < 0.0001)."),
            ("Knockout Phenocopy", "ISL1-/- hPSCs confirmed the prediction, eliminating SC-beta cells and expanding SC-EC cells."),
            ("Epistatic Rescue", "Lentiviral ISL1 overexpression was sufficient to repress SC-EC fate across all mutant backgrounds (WT, ISL1-/-, PDX1-/-, PAX6-/-), while beta induction required PDX1/PAX6 cooperation.")
        ],
        "Mutant-response data enable causal prediction of lineage switches, establishing the conceptual bridge to deep learning models."
    )

    # SLIDE 9: The Combinatorial Scaling Problem -> Bridge to pertTF
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "THE COMPUTATIONAL IMPERATIVE", "The Combinatorial Barrier: Why We Need Transferable Perturbation Models", "Physical screening cannot scale across the combinatorial space of ~20,000 human genes and multi-stage lineages")
    
    # 2 Comparison Cards
    c1 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tb1 = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "EXPERIMENTAL LIMITATIONS"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(197, 48, 48) # Red-ish
    p.space_after = Pt(12)
    p = tf1.add_paragraph()
    p.text = "• Massive Combinatorial Explosion: ~20,000 genes x 14 cell types x 5 stages = >1,000,000 experimental conditions."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    p = tf1.add_paragraph()
    p.text = "• High Cost & Fragility: Stem cell differentiation and primary human tissues are technically prohibitive for genome-scale Perturb-seq."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    p = tf1.add_paragraph()
    p.text = "• Context-Specificity: Biological responses are highly non-linear and depend heavily on target cell developmental state."
    p.font.size = Pt(13)

    c2 = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9))
    c2.fill.solid()
    c2.fill.fore_color.rgb = TAKEAWAY_BG
    c2.line.color.rgb = SLATE_ACCENT
    c2.line.width = Pt(1.5)
    tb2 = s9.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "THE pertTF COMPUTATIONAL OBJECTIVE"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    p = tf2.add_paragraph()
    p.text = "• Learn Transferable Rules: Train a context-aware transformer on multi-stage knockout village data to predict perturbation effects in unmeasured contexts."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    p = tf2.add_paragraph()
    p.text = "• Predict Higher-Order Phenotypes: Go beyond point-wise gene expression to predict cell identity shifts and population composition changes (lochNESS)."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    p = tf2.add_paragraph()
    p.text = "• Enable Virtual Genetic Screens: Perform genome-wide in silico CRISPR screens and transfer predictions to primary human patient islets."
    p.font.size = Pt(13)

    # =========================================================================
    # SECTION 2: pertTF (Slides 10-22)
    # =========================================================================

    # SLIDE 10: pertTF Architecture & Multi-Task Loss
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "pertTF AI FRAMEWORK", "pertTF: A Context-Aware Transformer for Genetic Perturbation Prediction", "Ingesting single-cell transcriptomes and perturbation tokens to model non-linear gene regulatory shifts")
    add_fitted_image(s10, "clean_sources/pertTF_pdf_pages/page-04.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s10, 8.6, 1.8, 4.0, 5.0,
        "Model Formulation & Innovations",
        [
            ("Tokenization & Input", "Inputs concatenated gene identity (Eg) and log-expression (Ex) tokens, preceded by an empty special token (Ec) aggregating whole-cell context."),
            ("Perturbation Integration", "Learnable perturbation adapter integrates genetic knockout identity directly into cell embeddings."),
            ("Negative Binomial Loss", "Replaces MSE with Negative Binomial NLL loss to capture single-cell count distributions, sampling 2:1 HVGs vs non-HVGs."),
            ("Contrastive Latent Space", "Supervised contrastive loss explicitly maximizes separation between disparate cell types and genotypes.")
        ],
        "pertTF learns a unified latent space co-embedding cell developmental state and perturbation-driven regulatory reconfigurations."
    )

    # SLIDE 11: Internal Representation & Baseline Classification Benchmark
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "pertTF AI FRAMEWORK", "Multi-Task Learning Yields High-Resolution Cellular and Genotype Embeddings", "pertTF decisively outperforms single-cell foundation models in cell-type and masked perturbation classification")
    add_fitted_image(s11, "clean_sources/pertTF_extracted_pptx_images/Fig._2_v5_slide1_img3.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s11, 8.6, 1.8, 4.0, 5.0,
        "Representation Benchmarks",
        [
            ("Cell-Type Classification", "The 12-layer pertTF model achieved Macro F1 > 0.98, Accuracy > 0.98 (vs scGPT F1 ~ 0.93, scFoundation F1 ~ 0.92, Geneformer F1 ~ 0.90)."),
            ("Masked Genotype Classification", "Predicting hidden perturbation status from expression alone, pertTF achieved Macro F1 ~ 0.84 (vs scGPT F1 ~ 0.61)."),
            ("Latent Cluster Separation", "Embeddings cleanly segregate the 14 cell lineages while simultaneously clustering genotypes within each stage (Fig. 1e-f).")
        ],
        "pertTF captures fine-grained transcriptomic state distinctions without compressing away perturbation-specific signatures."
    )

    # SLIDE 12: Higher-Order Phenotype: lochNESS Prediction
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "pertTF AI FRAMEWORK", "Beyond Expression: lochNESS Quantifies Higher-Order Cell Composition Shifts", "Continuous local density scoring bridges single-cell latent representations to population-level lineage diverticula")
    add_fitted_image(s12, "clean_sources/pertTF_extracted_pptx_images/Fig._2_v5_slide1_img0.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s12, 8.6, 1.8, 4.0, 5.0,
        "lochNESS Formulation & Validation",
        [
            ("lochNESS Concept", "Quantifies single-cell k-NN density ratio of perturbed vs WT cells: positive lochNESS (>0) denotes local trapping/enrichment; negative (<0) denotes depletion."),
            ("High Prediction Fidelity", "pertTF predicted lochNESS accurately across diverse knockouts (Pearson r ~ 0.88 for PDX1, r ~ 0.84 for TADA2B, r ~ 0.82 for GATA6)."),
            ("Superiority Over Expression", "Direct lochNESS regression achieved ROC-AUC ~ 0.86 in identifying enriched/depleted lineages (vs AUC ~ 0.64 for expression baseline).")
        ],
        "Latent manifold geometry captures cell fate redirection far better than point-wise differential gene expression averages."
    )

    # SLIDE 13: Unseen Cell Context Generalization
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "pertTF AI FRAMEWORK", "Predicting Perturbation Outcomes in Held-Out, Unseen Cellular Contexts", "pertTF infers the consequence of knocking out lineage directors in cell types never seen in perturbed form during training")
    add_fitted_image(s13, "clean_sources/pertTF_pdf_pages/page-20.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s13, 8.6, 1.8, 4.0, 5.0,
        "Unseen Context Experimental Design",
        [
            ("Held-Out Evaluation", "Perturbed cells of a target lineage (e.g. PDX1 KO in mature SC-beta cells) were completely withheld during training. The model only saw WT SC-beta and PDX1 KO in earlier stages."),
            ("Inference Task", "Provide WT SC-beta cells + PDX1 perturbation token; predict single-cell embedding shift and differential expression."),
            ("Benchmark Performance", "pertTF achieved Cosine Similarity ~ 0.91, PCC-delta ~ 0.72, and DE-direction matching ~ 85% (significantly exceeding scGPT Cosine ~ 0.74, GEARS ~ 0.68).")
        ],
        "pertTF learns transferable regulatory logic, understanding lineage-specific vulnerability without memorizing training pairs."
    )

    # SLIDE 14: Leave-One-Genotype-Out Testing (Unseen Genes)
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "pertTF AI FRAMEWORK", "Leave-One-Genotype-Out Testing Demonstrates Generalization to Unseen Genes", "Gene regulatory graph embeddings enable accurate prediction of knockouts entirely absent from training data")
    add_fitted_image(s14, "clean_sources/pertTF_extracted_pptx_images/Fig._2_v5_slide1_img4.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s14, 8.6, 1.8, 4.0, 5.0,
        "Unseen Gene Generalization",
        [
            ("GNN Knowledge Graph Prior", "Gene Ontology and protein interaction graph embeddings provide functional prior vectors for unseen genes."),
            ("Systematic LOGO Benchmark", "30 independent leave-one-genotype-out models were trained and tested across all 30 knockout lines."),
            ("Benchmark Winner", "pertTF achieved mean Cosine Similarity ~ 0.86 across held-out genes, outperforming scGPT (mean ~ 0.70) and GEARS (~ 0.64) in 27 out of 30 genotypes.")
        ],
        "pertTF successfully extrapolates perturbation effects to novel genes outside the experimental training set."
    )

    # SLIDE 15: Joint Generalization (Unseen Gene + Unseen Context)
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "pertTF AI FRAMEWORK", "Joint Generalization: Predicting Unseen Genes in Unseen Cell Contexts", "Solving the hardest out-of-distribution challenge: predicting novel knockouts in unmeasured cellular lineages")
    add_fitted_image(s15, "clean_sources/pertTF_extracted_pptx_images/Supplementary_Fig_v8_slide4_img24.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s15, 8.6, 1.8, 4.0, 5.0,
        "Joint Out-of-Distribution Regime",
        [
            ("Simultaneous Extrapolation", "Predicts the transcriptomic shift of an unseen gene knockout in a lineage where no perturbed cells were ever seen during training."),
            ("Quantitative Benchmark", "pertTF achieved Cosine Similarity ~ 0.81 and PCC-delta ~ 0.61 (vs scGPT Cosine ~ 0.59, PCC-delta ~ 0.34)."),
            ("Biological Implication", "Establishes pertTF as a robust generalizer capable of extrapolating across both biological axes simultaneously.")
        ],
        "pertTF maintains stable predictive performance under compound out-of-distribution evaluation regimes."
    )

    # SLIDE 16: Orthogonal Validation: CRISPRi Perturb-seq
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "pertTF AI FRAMEWORK", "Orthogonal Validation: Generalizing to Partial CRISPRi Repression", "Validation across an independent 50-gene chromatin/TF CRISPRi screen in human pluripotent stem cells")
    add_fitted_image(s16, "clean_sources/pertTF_extracted_pptx_images/Supplementary_Fig_v8_slide5_img26.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s16, 8.6, 1.8, 4.0, 5.0,
        "Cross-Modality Transfer",
        [
            ("Orthogonal Perturbation Modality", "CRISPRi introduces partial, continuous knockdown and variable penetrance compared to full genomic null alleles."),
            ("Isolating Responders", "Mixscape classification and continuous PS scoring were utilized to identify true responding cells ('KO') versus escapers ('NP')."),
            ("CTNNB1 Shift Prediction", "pertTF accurately predicted the massive transcriptomic shift of CTNNB1 knockdown (Cosine Similarity ~ 0.89 vs scGPT ~ 0.68; DEG Pearson r ~ 0.78).")
        ],
        "Learned perturbation representations reflect universal regulatory dynamics invariant to the delivery mechanism (Cas9 vs dCas9-KRAB)."
    )

    # SLIDE 17: Clinical Translation: Primary Human Islets
    s17 = prs.slides.add_slide(blank_layout)
    add_header(s17, "pertTF AI FRAMEWORK", "Clinical Transfer: Fine-Tuning pertTF on Primary Human Adult Islets", "Adapting in vitro stem cell-derived representations to primary donor tissue with fewer than 300 cells")
    add_fitted_image(s17, "clean_sources/pertTF_extracted_pptx_images/Fig._5_v4_slide1_img0.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s17, 8.6, 1.8, 4.0, 5.0,
        "Primary Islet Transfer Strategy",
        [
            ("The Clinical Challenge", "Primary human islets cannot be screened experimentally at scale due to limited donor material and post-mitotic fragility."),
            ("Rapid Fine-Tuning", "pertTF adapted to primary adult islet cell types (beta-1, beta-2, alpha, delta, ductal) with <300 primary cells within 5 epochs (F1 > 0.96)."),
            ("Latent State Inference", "Fine-tuned pertTF models were deployed to score latent transcription factor deficiency states across single cells from non-diabetic, pre-T2D, and T2D donors.")
        ],
        "Transfer learning allows in vitro perturbation knowledge to be projected onto inaccessible clinical patient tissues."
    )

    # SLIDE 18: Latent TF Disruption in Clinical T2D Donors
    s18 = prs.slides.add_slide(blank_layout)
    add_header(s18, "pertTF AI FRAMEWORK", "Latent PDX1, NEUROD1, and HNF4A Disruption Enriches in Clinical T2D Donors", "Clinical type 2 diabetes islet pathology converges onto developmental regulatory failure modes")
    add_fitted_image(s18, "clean_sources/pertTF_extracted_pptx_images/Supplementary_Fig_v8_slide6_img29.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s18, 8.6, 1.8, 4.0, 5.0,
        "Clinical Pathology Convergence",
        [
            ("PDX1 Loss in T2D", "Cells classified into a latent 'PDX1 loss state' increased by ~3.8-fold in T2D patient islets relative to non-diabetic controls (p < 0.0001)."),
            ("Beta-2 Subpopulation Vulnerability", "Fragile, disease-prone beta-2 cells showed marked depletion of WT state and significant enrichment of inferred NEUROD1 (p < 0.0001) and HNF4A (p < 0.001) loss states."),
            ("Dosage Concordance", "Single-cell HNF4A disruption probability strongly correlated with signature gene loss (PCC = -0.78, p = 0).")
        ],
        "In silico perturbation mapping reveals that adult T2D islet dysfunction mirrors the exact regulatory cascades identified in developmental knockouts."
    )

    # SLIDE 19: Primary Islet siRNA Experimental Validation
    s19 = prs.slides.add_slide(blank_layout)
    add_header(s19, "pertTF AI FRAMEWORK", "Primary Islet siRNA Validation Confirms In Silico Latent Predictions", "Experimental knockdown of RFX6 in primary human islets validates model classification accuracy")
    add_fitted_image(s19, "clean_sources/pertTF_extracted_pptx_images/Supplementary_Fig_v8_slide6_img31.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s19, 8.6, 1.8, 4.0, 5.0,
        "Experimental Primary Validation",
        [
            ("RFX6 siRNA Knockdown", "Primary human islets were treated with RFX6 siRNA and profiled via scRNA-seq to provide ground-truth perturbation data in primary tissue."),
            ("High Classification Accuracy", "pertTF accurately classified ~78% of RFX6 siRNA knockdown cells as RFX6-perturbed (versus <5% in non-targeting control cells)."),
            ("Manifold Alignment", "Predicted RFX6 perturbed cell embeddings aligned closely with observed primary siRNA embeddings (Cosine Similarity ~ 0.87).")
        ],
        "Experimental testing in primary human tissue confirms that pertTF accurately identifies authentic biological perturbation states."
    )

    # SLIDE 20: Virtual Genetic Screening (In Silico Screens)
    s20 = prs.slides.add_slide(blank_layout)
    add_header(s20, "pertTF AI FRAMEWORK", "Virtual Genetic Screening: In Silico Prioritization of Progenitor Regulators", "Genome-wide in silico screening accurately recovers known pancreatic lineage directors without physical screening")
    add_fitted_image(s20, "clean_sources/pertTF_extracted_pptx_images/Fig._6_v7_slide1_img4.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s20, 8.6, 1.8, 4.0, 5.0,
        "In Silico Screening Pipelines",
        [
            ("Method 1 (Embedding Similarity)", "Perturbs candidate genes in silico and computes cosine similarity to a target phenotype vector (e.g. PDX1- loss in Pancreatic Progenitors)."),
            ("Validation Against Physical Screens", "Benchmarked against an experimental pooled CRISPR screen sorting on PDX1-GFP (Wang et al., >18,000 genes)."),
            ("Decisive Benchmark Superiority", "pertTF achieved ROC-AUC = 0.79 (TEXT-VERIFIED) and AUPR ~ 0.74, markedly outperforming differential gene expression ranking (AUC = 0.66, AUPR ~ 0.52). Top hits recovered GATA6, MAPK1, PROX1.")
        ],
        "pertTF serves as a validated computational alternative to expensive pooled CRISPR screens for functional regulator discovery."
    )

    # SLIDE 21: Essential Gene Calibration & In Silico Perturb-seq
    s21 = prs.slides.add_slide(blank_layout)
    add_header(s21, "pertTF AI FRAMEWORK", "Essential-Gene Calibration and Genome-Scale In Silico Perturb-seq", "Predicting multi-protein complex dynamics and genome-wide single-cell expression profiles")
    add_fitted_image(s21, "clean_sources/pertTF_extracted_pptx_images/Fig._6_v7_slide1_img3.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s21, 8.6, 1.8, 4.0, 5.0,
        "Genome-Scale Extrapolation",
        [
            ("Essential Gene Prior Calibration", "Calibrating lochNESS with ribosomal and proteasomal essential genes accurately predicted negative lochNESS for unseen essential genes (e.g. MRPS5, p < 10^-15)."),
            ("Co-Clustering Multi-Protein Complexes", "In silico knockout of subunits within the same chromatin complexes (SMARCC1 / SMARCD1, SALL4 / TCF7L1) clustered together in embedding space."),
            ("Single-Cell Expression Profiles", "Predicted DEG profiles replicated observed Perturb-seq data with Pearson r > 0.81 across chromatin regulators.")
        ],
        "pertTF enables virtual single-cell Perturb-seq, capturing multi-subunit epistatic interactions computationally."
    )

    # SLIDE 22: pertTF Computational Synthesis
    s22 = prs.slides.add_slide(blank_layout)
    add_header(s22, "pertTF AI FRAMEWORK", "pertTF Computational Synthesis: Context-Aware Modeling Across Biological Scales", "A rigorous summary of verified model capabilities, biological discoveries, and conceptual boundaries")
    
    # 2 Summary Cards
    c1 = s22.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = TAKEAWAY_BG
    c1.line.color.rgb = SLATE_ACCENT
    c1.line.width = Pt(1.5)
    tb1 = s22.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "VERIFIED COMPUTATIONAL ADVANCES"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(12)
    p = tf1.add_paragraph()
    p.text = "• Generalizable Representations: Superior cell-type (F1 > 0.98) and genotype (F1 ~ 0.84) latent separability."
    p.font.size = Pt(12.5)
    p.space_after = Pt(8)
    p = tf1.add_paragraph()
    p.text = "• Higher-Order Phenotypes: Accurate lochNESS composition regression (r ~ 0.88, AUC ~ 0.86)."
    p.font.size = Pt(12.5)
    p.space_after = Pt(8)
    p = tf1.add_paragraph()
    p.text = "• Robust Out-of-Distribution Generalization: Excels in unseen context (Cosine ~ 0.91), unseen genes (mean ~ 0.86), and joint extrapolation (~ 0.81)."
    p.font.size = Pt(12.5)
    p.space_after = Pt(8)
    p = tf1.add_paragraph()
    p.text = "• Translational Impact: Bridges in vitro models to primary clinical T2D islets and genome-scale virtual screens (AUC = 0.79)."
    p.font.size = Pt(12.5)

    c2 = s22.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tb2 = s22.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "CONCEPTUAL BOUNDARIES & LIMITATIONS"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(197, 48, 48)
    p.space_after = Pt(12)
    p = tf2.add_paragraph()
    p.text = "• Prediction != Causal Proof: Inferred latent TF loss indicates transcriptomic similarity, not primary genomic lesion."
    p.font.size = Pt(12.5)
    p.space_after = Pt(8)
    p = tf2.add_paragraph()
    p.text = "• Embedding Compression: End-to-end neural network embeddings compress multiple distinct biological phenomena into single latent vectors."
    p.font.size = Pt(12.5)
    p.space_after = Pt(8)
    p = tf2.add_paragraph()
    p.text = "• Unresolved Phenotypic Dimensions: Prediction alone does not tell us whether an effect is driven by high penetrance, massive global magnitude, or focal state trapping."
    p.font.size = Pt(12.5)
    p.space_after = Pt(8)
    p = tf2.add_paragraph()
    p.text = "• Motivation for WIP: Decompose perturbation phenotypes into explicit, interpretable coordinates."
    p.font.size = Pt(12.5)

    # =========================================================================
    # SECTION 3: MULTI-DIMENSIONAL PHENOTYPIC DECOMPOSITION (WIP) (Slides 23-28)
    # =========================================================================

    # SLIDE 23: The Next Question -> Bridge to WIP
    s23 = prs.slides.add_slide(blank_layout)
    add_header(s23, "MULTI-DIMENSIONAL PHENOTYPING", "The Next Question: What Remains Compressed Inside a Perturbation Prediction?", "Deconstructing end-to-end latent predictions into explicit, interpretable biological coordinates")
    
    c1 = s23.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tb1 = s23.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11.1), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "THE MOTIVATION FOR MULTI-SCALE DECOMPOSITION"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = CYAN_DARK
    p.space_after = Pt(14)
    
    p = tf1.add_paragraph()
    p.text = "pertTF provides state-of-the-art predictive accuracy across diverse unseen contexts and virtual screens. However, predicting a shifted latent embedding compresses multiple independent biological questions:"
    p.font.size = Pt(13)
    p.space_after = Pt(12)
    
    p = tf1.add_paragraph()
    p.text = "1. Single-Cell Penetrance: What proportion of single cells carrying the mutation genuinely depart baseline? (PS)"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(8)
    
    p = tf1.add_paragraph()
    p.text = "2. Global Multivariate Magnitude: How far is the overall population displaced in multivariate transcriptomic space? (Energy Distance)"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(8)
    
    p = tf1.add_paragraph()
    p.text = "3. Directional Topological Localization: Where on the single-cell manifold do cells specifically accumulate or drop out? (Signed lochNESS)"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(14)
    
    p = tf1.add_paragraph()
    p.text = "The WIP establishes a multi-dimensional framework to evaluate whether these properties are coupled or decoupled across human developmental regulators."
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = MUTED_GRAY

    # SLIDE 24: Three Complementary Phenotypic Dimensions
    s24 = prs.slides.add_slide(blank_layout)
    add_header(s24, "MULTI-DIMENSIONAL PHENOTYPING", "Three Complementary Dimensions of Single-Cell Perturbation Phenotypes", "Formalizing response penetrance (PS), multivariate magnitude (Energy Distance), and state localization (lochNESS)")
    add_fitted_image(s24, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/19_perturbation_summary.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s24, 8.6, 1.8, 4.0, 5.0,
        "Three Phenotypic Pillars",
        [
            ("PS (Single-Cell Penetrance)", "Continuous [0, 1] projection score via PS_python measuring single-cell response probability. Summarized via median PS and responder fraction (PS > 0.5). GLIS3 exhibits #1 highest PS (0.750, 66.6% responders)."),
            ("Energy Distance (Magnitude)", "Multivariate statistical displacement in PCA space vs WT. All 36 tested non-WT lines show significant divergence (FDR = 0.000999, permutation floor). PDX1het exhibits #1 highest Energy Distance (13.6337, MMD 0.3129)."),
            ("Signed lochNESS (Localization)", "Signed k-NN local density ratio resolving positive enrichment (trapping) from negative depletion (loss). Peak: FOXA2 in Liver (+13.740, Peak cell +18.307).")
        ],
        "These three coordinates define a multi-dimensional phenotype space capturing distinct facets of genetic disruption."
    )

    # SLIDE 25: Coupling: PS vs Energy Distance
    s25 = prs.slides.add_slide(blank_layout)
    add_header(s25, "MULTI-DIMENSIONAL PHENOTYPING", "Response Penetrance Strongly Couples with Global Multivariate Magnitude", "Perturbations with high single-cell response rates produce proportionally larger population-level displacements")
    add_fitted_image(s25, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/07_ps_vs_distance.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s25, 8.6, 1.8, 4.0, 5.0,
        "Strong Cross-Metric Coupling",
        [
            ("Statistical Association", "Median PS and Energy Distance exhibit strong, highly significant positive correlation: Spearman rho = +0.6410, p = 4.18 x 10^-4 (Pearson r = +0.7628, p = 5.87 x 10^-6, n = 26)."),
            ("High-Impact Regulators", "Factors like GLIS3, GATA6, HHEX, and KDM2B combine high single-cell penetrance (PS > 0.60) with massive global transcriptomic divergence (Energy Distance > 7.9)."),
            ("Biological Meaning", "A high penetrance rate across single cells directly translates into a large statistical displacement of the population distribution.")
        ],
        "Single-cell response strength is a primary determinant of overall multivariate transcriptomic divergence."
    )

    # SLIDE 26: Decoupling: PS vs Directional lochNESS
    s26 = prs.slides.add_slide(blank_layout)
    add_header(s26, "MULTI-DIMENSIONAL PHENOTYPING", "Single-Cell Penetrance Decouples from Directional State Localization", "A strong single-cell response rate does not dictate where mutant cells localize on the developmental manifold")
    add_fitted_image(s26, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/10_ps_vs_lochness_positive.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s26, 8.6, 1.8, 4.0, 5.0,
        "Decoupling Analysis Across Dimensions",
        [
            ("PS vs Positive lochNESS", "Completely decoupled: Spearman rho = +0.0612, p = 0.7665 (n = 26). High penetrance does not imply strong focal trapping."),
            ("PS vs Negative lochNESS", "Completely decoupled: Spearman rho = -0.1715, p = 0.4123 (n = 25). High penetrance does not dictate lineage dropout severity."),
            ("Energy Distance vs lochNESS", "Moderately coupled with absolute lochNESS (rho = +0.4512, p = 0.0057) and positive lochNESS (rho = +0.4103, p = 0.0129), but uncoupled from negative lochNESS (rho = -0.1410, p = 0.4263).")
        ],
        "PS, Energy Distance, and signed lochNESS provide complementary, non-redundant coordinates of perturbation biology."
    )

    # SLIDE 27: Integrated Case Studies: PDX1, GATA6, and FOXA2
    s27 = prs.slides.add_slide(blank_layout)
    add_header(s27, "CASE STUDY SYNTHESIS", "Multi-Scale Case Studies: Tracking PDX1, GATA6, and FOXA2 Across All Stages", "Traversing biological phenotype, pertTF model prediction, and WIP multi-dimensional decomposition")
    add_fitted_image(s27, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/28_umap_highlight_genotypes.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s27, 8.6, 1.8, 4.0, 5.0,
        "Audited Case Study Profiles",
        [
            ("PDX1 / PDX1het (Master Beta Selector)", "Precursor: SC-beta collapse & SC-EC trade-off. pertTF: lochNESS r ~ 0.88; held-out SC-beta Cosine ~ 0.91; T2D donor enrichment ~ 3.8-fold. WIP: PDX1het has #1 highest Energy Distance (13.6337, MMD 0.3129, PG2); PDX1 hom has Energy Dist = 2.4579 (PG7, Module M4)."),
            ("GATA6 (Endoderm Gatekeeper)", "Precursor: Endoderm failure & Endothelial diversion. pertTF: Top PP virtual screen hit. WIP: Median PS = 0.6482 (61.6% resp), Energy Dist = 9.7045 (MMD 0.1590), Endothelial lochNESS = +4.530 (Peak +37.059), Module M1."),
            ("FOXA2 (Foregut Pioneer)", "Precursor: Foregut failure & Hepatic diversion. pertTF: Pioneer network capture. WIP: Energy Dist = 5.4466 (MMD 0.1063), dominant Liver lochNESS = +13.740 (Peak cell +18.307), Module M1.")
        ],
        "Rigorous cross-stage tracking grounds computational predictions in validated developmental biology."
    )

    # SLIDE 28: DistanceSpace vs Stage 7 Modules
    s28 = prs.slides.add_slide(blank_layout)
    add_header(s28, "MULTI-DIMENSIONAL PHENOTYPING", "DistanceSpace and Downstream Modules Reveal Higher-Order Architecture", "PCoA phenotypic similarity geometry and perturbation-to-gene regulatory effect modules")
    add_fitted_image(s28, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/14_distance_space.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s28, 8.6, 1.8, 4.0, 5.0,
        "Two Distinct Analytical Dimensions",
        [
            ("DistanceSpace (Pairwise Geometry)", "Computes 630 pairwise Energy Distances across 36 non-WT lines, partitioning perturbations into 9 Phenotype Groups (PG1-PG9) via PCoA/neighborhood similarity."),
            ("Stage 7 Regulatory Modules", "Analyzes the perturbation x downstream-gene effect matrix, clustering into 6 Co-Functional Modules (M1: Core Endoderm; M3: NEUROD1/BCOR; M4: PDX1; M6: PAX6/PBX1)."),
            ("Core Gene Programs", "cNMF identifies 4 downstream programs (P1: Secretion; P2: Hypoxia/Survival; P3/P4: Proliferation/ECM).")
        ],
        "DistanceSpace maps phenotypic similarity geometry, while Stage 7 uncovers shared downstream regulatory targets."
    )

    # =========================================================================
    # SECTION 4: CONCLUSION & FUTURE DIRECTIONS (Slides 29-30)
    # =========================================================================

    # SLIDE 29: Master Summary: Biology to AI
    s29 = prs.slides.add_slide(blank_layout)
    add_header(s29, "MASTER SCIENTIFIC SYNTHESIS", "A Complete Scientific Trajectory: From Developmental Genetics to AI", "Integrating causal stem cell perturbation biology with context-aware AI modeling and multi-dimensional phenotyping")
    
    col_w = 3.65
    col_gap = 0.38
    col_top = 1.8
    col_h = 4.9
    
    c1 = s29.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(col_top), Inches(col_w), Inches(col_h))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tb1 = s29.shapes.add_textbox(Inches(1.0), Inches(col_top + 0.2), Inches(col_w - 0.4), Inches(col_h - 0.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "1. BIOLOGY DISCOVERY"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = CYAN_DARK
    p.space_after = Pt(8)
    p = tf1.add_paragraph()
    p.text = "• Longitudinal Knockout Village: Profiled 30 genes across 5 stages (>111,000 cells)."
    p.font.size = Pt(12)
    p.space_after = Pt(6)
    p = tf1.add_paragraph()
    p.text = "• Lineage Rewiring: Monogenic diabetes mutations divert progenitors into SC-EC enterochromaffin cells."
    p.font.size = Pt(12)
    p.space_after = Pt(6)
    p = tf1.add_paragraph()
    p.text = "• ISL1 Repressor: Predictive modeling identified ISL1 as the master gatekeeper safeguarding beta fate."
    p.font.size = Pt(12)

    c2 = s29.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + col_w + col_gap), Inches(col_top), Inches(col_w), Inches(col_h))
    c2.fill.solid()
    c2.fill.fore_color.rgb = TAKEAWAY_BG
    c2.line.color.rgb = SLATE_ACCENT
    c2.line.width = Pt(1.5)
    tb2 = s29.shapes.add_textbox(Inches(1.0 + col_w + col_gap), Inches(col_top + 0.2), Inches(col_w - 0.4), Inches(col_h - 0.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "2. pertTF AI MODEL"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(8)
    p = tf2.add_paragraph()
    p.text = "• Context-Aware Architecture: Transformer with NB-NLL loss and supervised contrastive latent learning."
    p.font.size = Pt(12)
    p.space_after = Pt(6)
    p = tf2.add_paragraph()
    p.text = "• lochNESS Prediction: Captures population composition shifts beyond point-wise gene expression."
    p.font.size = Pt(12)
    p.space_after = Pt(6)
    p = tf2.add_paragraph()
    p.text = "• Out-of-Distribution Generalization: Excels in unseen context, unseen genes, primary T2D islets, and virtual screens."
    p.font.size = Pt(12)

    c3 = s29.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8 + (col_w + col_gap)*2), Inches(col_top), Inches(col_w), Inches(col_h))
    c3.fill.solid()
    c3.fill.fore_color.rgb = CARD_BG
    c3.line.color.rgb = CARD_BORDER
    tb3 = s29.shapes.add_textbox(Inches(1.0 + (col_w + col_gap)*2), Inches(col_top + 0.2), Inches(col_w - 0.4), Inches(col_h - 0.4))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "3. MULTI-SCALE PHENOTYPING"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = CYAN_DARK
    p.space_after = Pt(8)
    p = tf3.add_paragraph()
    p.text = "• Decompression: Deconstructs perturbations into PS penetrance, Energy Distance magnitude, and signed lochNESS localization."
    p.font.size = Pt(12)
    p.space_after = Pt(6)
    p = tf3.add_paragraph()
    p.text = "• Decoupling: PS couples with Energy Distance (rho = +0.641), but decouples from state localization."
    p.font.size = Pt(12)
    p.space_after = Pt(6)
    p = tf3.add_paragraph()
    p.text = "• Multi-Scale Mapping: Provides interpretable coordinates for next-generation foundation models."
    p.font.size = Pt(12)

    # SLIDE 30: Future Modeling Outlook
    s30 = prs.slides.add_slide(blank_layout)
    add_header(s30, "FUTURE MODELING OUTLOOK", "Toward Next-Generation Phenotype-Aware Perturbation Models", "Integrating multi-dimensional coordinates directly into neural network loss functions")
    
    c1 = s30.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tb1 = s30.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11.1), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "ROADMAP FOR PREDICTIVE BIOLOGY & FOUNDATION MODELS"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(14)
    
    p = tf1.add_paragraph()
    p.text = "1. Multi-Objective Loss Engineering: Incorporating Energy Distance and signed lochNESS directly into transformer pre-training loss functions to force models to learn population-level distributional shifts explicitly."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    
    p = tf1.add_paragraph()
    p.text = "2. Scaled Perturb-seq Ingestion: Expanding pertTF training across emerging massive datasets (e.g. Tahoe-100M, MorPhiC consortium null alleles) spanning thousands of genes across human lineages."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    
    p = tf1.add_paragraph()
    p.text = "3. Multimodal Integration: Incorporating single-cell chromatin accessibility (scATAC-seq), spatial transcriptomics, and protein abundance to model multi-omic epigenomic constraints."
    p.font.size = Pt(13)
    p.space_after = Pt(10)
    
    p = tf1.add_paragraph()
    p.text = "4. Therapeutic In Silico Screening: Deploying fine-tuned models to screen combination genetic perturbations and pharmacological interventions for regenerative cell therapies in diabetes."
    p.font.size = Pt(13)

    # =========================================================================
    # SECTION 5: BACKUP SLIDES (Slides 31-34)
    # =========================================================================

    # SLIDE 31: Backup | DistanceTest Methodology & Resolution Floor
    s31 = prs.slides.add_slide(blank_layout)
    add_header(s31, "BACKUP | TECHNICAL METHODOLOGY", "DistanceTest Methodology and Permutation Resolution Floor", "Statistical framework for multivariate two-sample Energy Distance permutation testing")
    add_fitted_image(s31, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/06_energy_distance_by_perturbation.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s31, 8.6, 1.8, 4.0, 5.0,
        "DistanceTest Audit",
        [
            ("Mathematical Formulation", "Calculated in 50-dimensional PCA space against a bounded control sample of N=5,000 WT cells."),
            ("Permutation Resolution Floor", "1,000 permutations impose an empirical p-value resolution limit of 1/(N_perm+1) ~ 0.000999. All 36 tested non-WT lines achieved FDR = 0.000999."),
            ("Effect Magnitude Distinction", "The identical permutation-bounded FDR values reflect this statistical resolution limit. Biological effect size must be evaluated from the Energy Distance (ranging from 1.64 to 13.63).")
        ],
        "Energy Distance quantifies the true biological displacement, while permutation testing establishes statistical divergence."
    )

    # SLIDE 32: Backup | Single-Cell PS vs lochNESS Atlases
    s32 = prs.slides.add_slide(blank_layout)
    add_header(s32, "BACKUP | FULL MANIFOLD ATLASES", "Single-Cell PS and lochNESS Manifold Atlases Across 111,581 Cells", "Global distribution of response penetrance and local neighborhood enrichment on UMAP")
    add_fitted_image(s32, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/27_umap_ps_lochness_comparison.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s32, 8.6, 1.8, 4.0, 5.0,
        "Full Manifold Visualization",
        [
            ("Spatial PS Distribution", "High PS scores concentrate in specific responsive differentiation states while unperturbed cells and escapers remain near baseline."),
            ("Signed lochNESS Polarity", "Resolves positive trapping (>0) from negative lineage depletion (<0) across the entire single-cell manifold."),
            ("Atlas Utility", "Provides an exhaustive single-cell atlas benchmark for all 37 genotypes.")
        ],
        "Manifold-wide visualization confirms the complementary distribution of PS penetrance and signed lochNESS localization."
    )

    # SLIDE 33: Backup | High-Resolution Genotype x Cell-State Response Matrices
    s33 = prs.slides.add_slide(blank_layout)
    add_header(s33, "BACKUP | RESPONSE MATRICES", "High-Resolution Genotype x Cell-State Response Matrices", "Deconvolving single-cell penetrance (PS) and directional localization across all 14 cell types")
    add_fitted_image(s33, "/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures/13_lochness_by_celltype.png", 0.8, 1.8, 7.5, 5.0)
    add_interpretation_card(s33, 8.6, 1.8, 4.0, 5.0,
        "High-Resolution Heatmaps",
        [
            ("State-Specific Penetrance", "PS response distributes selectively across developmental checkpoints rather than uniformly across lineages."),
            ("Bipolar lochNESS Profiles", "Captures reciprocal transitions: simultaneous severe depletion in mature SC-beta and enrichment in EnP/SC-EC (e.g. PDX1, RFX6, PAX6)."),
            ("Lineage Specificity Matrix", "Provides an exhaustive reference table for all 37 genotypes across 14 curated cell lineages.")
        ],
        "Lineage-resolved response heatmaps expose the exact developmental stages where genetic perturbations act."
    )

    # SLIDE 34: Backup | Methodological Audit & Skipped Targets
    s34 = prs.slides.add_slide(blank_layout)
    add_header(s34, "BACKUP | METHODOLOGICAL AUDIT", "Quality Control, Bounded Sampling, and Skipped Targets Audit", "Technical audit of sampling thresholds, missing data handling, and computational pipelines")
    
    c1 = s34.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.9))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tb1 = s34.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11.1), Inches(4.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "TECHNICAL QUALITY CONTROL & PIPELINE AUDIT"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = CYAN_DARK
    p.space_after = Pt(14)
    
    p = tf1.add_paragraph()
    p.text = "• Bounded WT Sampling: Control distributions were subsampled to N=5,000 cells to prevent sampling bias while preserving high statistical power for multivariate two-sample testing."
    p.font.size = Pt(12.5)
    p.space_after = Pt(10)
    
    p = tf1.add_paragraph()
    p.text = "• Skipped PS Targets Clarification: Exactly 10 genotypes (GATA4het, GATA6het, HHEXe, HHEXhet, HNF4Ahet, NANOGe-het, ONECUT1e, PDX1het, QSER1TET1, TET1/2/3) were skipped in PS calculation because their compound identifiers do not match an individual single gene symbol in the RNA count matrix. This is strictly an analytical lookup constraint, NOT zero biological response (e.g. PDX1het has the highest Energy Distance in the entire dataset: 13.6337)."
    p.font.size = Pt(12.5)
    p.space_after = Pt(10)
    
    p = tf1.add_paragraph()
    p.text = "• Allele Label Integrity: All heterozygous lines (e.g. PDX1het, GATA6het, HHEXhet) and enhancer deletions (e.g. HHEXe, ONECUT1e) are maintained as distinct experimental perturbations and not conflated with homozygous knockouts."
    p.font.size = Pt(12.5)

    # Save presentation
    out_pptx = "PertTF_Master_Scientific_Story.pptx"
    prs.save(out_pptx)
    print(f"Successfully generated master deck: {out_pptx} ({len(prs.slides)} slides)")

if __name__ == '__main__':
    create_deck()
