#!/usr/bin/env python3
"""Build polished, auditable scientific PowerPoint presentation for Diabetes Perturb-seq Analysis.

Generates:
    results/diabetes_specific/presentation/diabetes_perturbseq_comprehensive_analysis.pptx
"""

import os
import sys
from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# -----------------------------------------------------------------------------
# Color Palette & Typography
# -----------------------------------------------------------------------------
COLOR_BG = RGBColor(255, 255, 255)         # Pure white
COLOR_DARK = RGBColor(26, 32, 44)          # Slate 900
COLOR_MUTED = RGBColor(74, 85, 104)        # Slate 600
COLOR_LIGHT_GRAY = RGBColor(247, 250, 252) # Slate 50
COLOR_BORDER = RGBColor(226, 232, 240)     # Slate 200

COLOR_PRIMARY = RGBColor(43, 108, 176)     # Deep Blue (#2B6CB0)
COLOR_SECONDARY = RGBColor(49, 151, 149)   # Teal (#319795)
COLOR_ACCENT = RGBColor(214, 158, 46)      # Amber / Gold (#D69E2E)
COLOR_ALERT = RGBColor(229, 62, 62)        # Crimson / Red (#E53E3E)
COLOR_GREEN = RGBColor(47, 133, 90)        # Forest Green (#2F855A)
COLOR_PURPLE = RGBColor(128, 90, 213)      # Violet (#805AD5)

FONT_HEADING = "Helvetica"
FONT_BODY = "Arial"

PRESENTATION_DIR = Path("/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/presentation")
FIGURES_DIR = Path("/local/projects-t3/lilab/vmenon/PertTF-Virtual-Challeng-Weilab/perturbseq-pipeline/results/diabetes_specific/figures")
PPTX_PATH = PRESENTATION_DIR / "diabetes_perturbseq_comprehensive_analysis.pptx"

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_header(slide, title_text, category_text="PANCREATIC PERTURB-SEQ ANALYSIS", slide_num=None, is_backup=False):
    """Add standardized clean header banner."""
    # Category / Super-title
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.5), Inches(0.3))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"BACKUP | {category_text}" if is_backup else category_text
    p_cat.font.name = FONT_HEADING
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_SECONDARY if not is_backup else COLOR_PURPLE

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.2), Inches(0.75))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_DARK

    # Slide Number & Footer
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf_foot = footer_box.text_frame
    tf_foot.word_wrap = True
    tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
    p_foot = tf_foot.paragraphs[0]
    num_str = f"Slide {slide_num} of 30" if slide_num else ""
    p_foot.text = f"Human Pancreatic Differentiation Perturb-seq (GSE216909) | 111,581 Cells x 36 TFs                                           {num_str}"
    p_foot.font.name = FONT_BODY
    p_foot.font.size = Pt(9)
    p_foot.font.color.rgb = RGBColor(160, 174, 192)

def add_what_result_why_card(slide, left, top, width, height, what_text, result_text, why_text):
    """Add standardized WHAT / RESULT / WHY structured callout card."""
    # Outer Card Background
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_LIGHT_GRAY
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1)

    # Content Textbox
    pad = Inches(0.12)
    tb = slide.shapes.add_textbox(left + pad, top + pad, width - (2 * pad), height - (2 * pad))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # WHAT Section
    p1 = tf.paragraphs[0]
    r1_tag = p1.add_run()
    r1_tag.text = "WHAT: "
    r1_tag.font.name = FONT_HEADING
    r1_tag.font.size = Pt(10)
    r1_tag.font.bold = True
    r1_tag.font.color.rgb = COLOR_PRIMARY
    r1_txt = p1.add_run()
    r1_txt.text = what_text + "\n"
    r1_txt.font.name = FONT_BODY
    r1_txt.font.size = Pt(9.5)
    r1_txt.font.color.rgb = COLOR_DARK

    # RESULT Section
    p2 = tf.add_paragraph()
    r2_tag = p2.add_run()
    r2_tag.text = "RESULT: "
    r2_tag.font.name = FONT_HEADING
    r2_tag.font.size = Pt(10)
    r2_tag.font.bold = True
    r2_tag.font.color.rgb = COLOR_GREEN
    r2_txt = p2.add_run()
    r2_txt.text = result_text + "\n"
    r2_txt.font.name = FONT_BODY
    r2_txt.font.size = Pt(9.5)
    r2_txt.font.color.rgb = COLOR_DARK

    # WHY Section
    p3 = tf.add_paragraph()
    r3_tag = p3.add_run()
    r3_tag.text = "WHY: "
    r3_tag.font.name = FONT_HEADING
    r3_tag.font.size = Pt(10)
    r3_tag.font.bold = True
    r3_tag.font.color.rgb = COLOR_ALERT
    r3_txt = p3.add_run()
    r3_txt.text = why_text
    r3_txt.font.name = FONT_BODY
    r3_txt.font.size = Pt(9.5)
    r3_txt.font.color.rgb = COLOR_DARK

def add_bullet_card(slide, left, top, width, height, title, items, title_color=COLOR_PRIMARY):
    """Add a structured card containing bullets."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_LIGHT_GRAY
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1)

    pad = Inches(0.12)
    tb = slide.shapes.add_textbox(left + pad, top + pad, width - (2 * pad), height - (2 * pad))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = title_color
        p0.space_after = Pt(4)
        first_item = True
    else:
        first_item = False

    for idx, item in enumerate(items):
        if first_item and idx == 0:
            p = tf.add_paragraph()
        elif not title and idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        p.text = f"• {item}"
        p.font.name = FONT_BODY
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(3)

def add_image_fit(slide, img_path, left, top, max_w, max_h):
    """Add image preserving aspect ratio within bounding box."""
    if not os.path.exists(img_path):
        print(f"Warning: image missing {img_path}")
        return None

    with Image.open(img_path) as im:
        orig_w, orig_h = im.size

    aspect = orig_w / orig_h
    box_aspect = max_w / max_h

    if aspect > box_aspect:
        # Width-limited
        w = max_w
        h = max_w / aspect
        x = left
        y = top + (max_h - h) / 2
    else:
        # Height-limited
        h = max_h
        w = max_h * aspect
        y = top
        x = left + (max_w - w) / 2

    pic = slide.shapes.add_picture(str(img_path), x, y, w, h)
    return pic

def set_speaker_notes(slide, notes_text):
    """Set formatted speaker notes for slide."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

print("Helper library loaded successfully!")
